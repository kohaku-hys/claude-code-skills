"""pptx_check.py の試験（2026-10-15）。python-pptx でデッキを作り、情報量・1枚1メッセージの手がかり・動きの予算の判定を見る。
Run:  py -3 test_pptx_check.py   (exit 0 = pass)
"""
import json, os, subprocess, sys, tempfile
from pptx import Presentation
from pptx.util import Cm, Pt

S = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "pptx_check.py")
d = tempfile.mkdtemp()
ok = 0; ng = 0


def check(cond, label):
    global ok, ng
    if cond:
        ok += 1
    else:
        ng += 1
        print("NG:", label)


def tb(slide, x, y, text, pt, name=None):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(20), Cm(1.5))
    if name:
        box.name = name
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(pt)
    return box


prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
L = prs.slide_layouts[6]
# 1 すっきり：見出し＋短い1行＋出典
s = prs.slides.add_slide(L)
tb(s, 2, 1.5, "Short videos use 3 tricks on our brain", 39)
tb(s, 2, 8, "Random rewards", 28)
tb(s, 2, 17.5, "Source 1: LINE Research (2024)", 12)
tb(s, 31, 18, "1 / 5", 12)
# 2 ぎちぎち：読む場所14・語数多い・小さい字
s = prs.slides.add_slide(L)
tb(s, 2, 1.5, "Dopamine cells react to better than expected", 39)
for k in range(13):
    tb(s, 2, 4 + k * 1.0, f"label number {k} with several more words here", 18)
# 3 見出しが2つの主張・文が3本
s = prs.slides.add_slide(L)
tb(s, 2, 1.5, "A maybe reward excites dopamine cells, so it may boost wanting", 39)
for k in range(3):
    tb(s, 2, 6 + k * 2, f"This is a full sentence number {k} on the slide", 24)
# 4 小さすぎる字（日本語）
s = prs.slides.add_slide(L)
tb(s, 2, 1.5, "ドーパミンは予想との差に反応する", 40)
tb(s, 2, 8, "細かい注記の文字", 12)
# 5 参考文献（数えない）
s = prs.slides.add_slide(L)
tb(s, 2, 1.5, "References", 39)
for k in range(15):
    tb(s, 2, 3 + k * 0.9, f"Author {k} (2020). A long title of a paper that is cited here. Journal.", 16)
src = os.path.join(d, "deck.pptx")
prs.save(src)


def run(*args):
    r = subprocess.run([sys.executable, S, src, *args], capture_output=True, text=True, encoding="utf-8",
                       env=dict(os.environ, PYTHONIOENCODING="utf-8"))
    return r.returncode, r.stdout + r.stderr


rc, out = run("--json")
res = json.loads(out)
by = {r["stat"]["slide"]: r for r in res["slides"]}
codes = lambda n, lv=None: [i["code"] for i in by[n]["issues"] if lv is None or i["level"] == lv]
check(rc == 1, f"ERROR があれば exit 1 ({rc})")
check(codes(1) == [], f"1 すっきりした枚は何も出ない {by[1]['issues']}")
check(by[1]["stat"]["blocks"] == 2, f"1 出典・ページ番号は読む場所に数えない ({by[1]['stat']['blocks']})")
check("I1" in codes(2, "ERROR") and "I2" in codes(2, "ERROR"), f"2 読む場所14・語数で停止 {codes(2)}")
check("I3" in codes(2, "WARNING"), f"2 21pt 未満が多いと注意 {codes(2)}")
check("O1" in codes(3) and "O2" in codes(3), f"3 見出しのつなぎ・文3本 {codes(3)}")
check("I3" in codes(4, "ERROR"), f"4 12pt の本文は停止 {codes(4)}")
check(codes(5) == [], f"5 参考文献の枚は数えない {codes(5)}")
rc, out = run("--draft")
check(rc == 0 and "ERROR 0" in out, f"--draft は ERROR を下げて exit 0 ({out.strip()[-30:]})")
rc, out = run("--slides", "1")
check(rc == 0 and "2枚目" not in out, "--slides で枚を絞る")

print(f"ok={ok} ng={ng}")
sys.exit(1 if ng else 0)
