# -*- coding: utf-8 -*-
"""pptx_check.py — できあがった pptx を1枚ずつ調べる（情報量・1枚1メッセージの手がかり・動きの予算）。/slide-gen・/slides の道具箱。

なぜ存在するか：deck_doctor は ppt-master の SVG（原稿）しか見ない。pptx を直接直す経路（作者の環境では /slides）や、
motion.py で後から足した動きは、どの関門も通らずに本人へ渡っていた。2026-10-15 本人「情報量が多いのと
動きができすぎてる」「1つのスライドに1つ伝えたいことを載せるようにスキル改善」（以前のデッキの6・7枚目＝
読む場所が16・15か所、文が6本、クリック4回＋点の点滅。デッキ全体でも1クリックで6〜10個が動く枚があった）。

使い方（python 3.9 以上・python-pptx）：
  py -3 pptx_check.py <deck.pptx> [--slides 6,7] [--json] [--draft]
  exit 0 = ERROR なし ／ 1 = ERROR あり（--draft は ERROR を WARNING に下げて 0 で返す。途中の確認用。本人に渡す版には付けない）

検査（数字は deck_doctor の話す型と同じ。英語は語数で数える）：
  情報量
    I1 読む場所（文字の枠の数。見出しは何枠でも1と数える。出典・※・ページ番号は数えない）：9個以上で警告・13個以上で停止
    I2 文字の量：日本語は 100字超で警告・150字超で停止／英語は 30語超で警告・45語超で停止
    I3 本文の文字の大きさ：16.5pt（22px）未満が1字でもあれば停止・21pt（28px）未満が30字（英語6語）を超えたら警告
  1枚1メッセージ（機械では手がかりだけ。決めるのは slide-reader の「1文で言い直せるか」）
    O1 見出しが2つの主張をつないでいる（英語 ", so" ", which" ", and" ", but" "; "／日本語「ので、」「から、」「し、」「が、」「ため、」）：警告
    O2 本文に「文」（英語6語以上・日本語15字以上の枠）が3本以上：警告＝伝えたいことが複数ある疑い
  動きの予算（motion.py と同じ数字。motion.py の BUDGET はここから読む）
    A1 クリック：3回で警告・4回以上で停止
    A2 1回のクリックで動く物（図形の数）：3つで警告・4つ以上で停止
    A3 自動で繰り返す動き：1つで警告（飾りでないか）・2つ以上で停止
    A4 強調（拡大/収縮・回転）と軌跡：1枚3つ以上で警告
    ノートの【動きの作り方】に「・動きの予算を超えた理由：…」の行がある枚は、A1〜A3 の停止を警告に下げる（motion.py の "busy"）。
  参考文献の枚（見出しが References / 参考文献 / 出典 / Sources）と、文字の無い枚は I・O を数えない。
"""
from __future__ import annotations

import argparse
import json
import re
import sys

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"p": P_NS, "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

# 情報量（deck_doctor の TALK_* に合わせる）
BLOCKS_WARN, BLOCKS_STOP = 9, 13
JA_WARN, JA_STOP = 100, 150
EN_WARN, EN_STOP = 30, 45
FLOOR_PT, SMALL_PT, SMALL_JA, SMALL_EN = 16.5, 21.0, 30, 6
SENTENCE_EN, SENTENCE_JA, SENTENCES_WARN = 6, 15, 3
TITLE_PT = 30.0  # これ以上の文字を含む上の方の枠を見出しとみなす
# 動きの予算
BUDGET = {"clicks_warn": 3, "clicks_stop": 4, "targets_warn": 3, "targets_stop": 4,
          "loops_warn": 1, "loops_stop": 2, "moves_warn": 3}
BUSY_MARK = "動きの予算を超えた理由"

FOOT_RE = re.compile(r"^\s*(Source|Sources|Photo|Photos|Image|出典|※|写真|画像|\*)", re.I)
PAGE_RE = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
REF_TITLE_RE = re.compile(r"^\s*(References?|Sources|Bibliography|参考文献|出典|引用文献)\s*$", re.I)
EN_WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'’\-]*")
JA_CHAR = re.compile(r"[぀-ヿ㐀-鿿ｦ-ﾟ]")
JOIN_EN = re.compile(r",\s*(so|which|and|but)\b|;\s", re.I)
JOIN_JA = re.compile(r"(ので|から|し|が|ため)、")


# --------------------------------------------------------------------------- 文字

def _walk(shapes, foot=False):
    for sh in shapes:
        is_foot = foot or "footer" in (sh.name or "").lower()
        if sh.shape_type == 6:  # group
            yield from _walk(sh.shapes, is_foot)
        else:
            yield sh, is_foot


def _sizes(tf):
    out = []
    for para in tf.paragraphs:
        for r in para.runs:
            if r.text.strip():
                out.append((r.font.size.pt if r.font.size is not None else None, r.text))
    return out


def measure(text):
    """(日本語の字数, 英語の語数)"""
    return len(JA_CHAR.findall(text)), len(EN_WORD.findall(text))


def is_ja(text):
    ja, en = measure(text)
    return ja > en


def text_items(slide, slide_h):
    """[(kind, text, sizes, top_cm)] kind = title / body / foot"""
    items = []
    for sh, foot in _walk(slide.shapes):
        if not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
            continue
        text = sh.text_frame.text.strip()
        if not text:
            continue
        sizes = _sizes(sh.text_frame)
        top = (sh.top or 0) / 360000
        if foot or FOOT_RE.match(text) or PAGE_RE.match(text):
            kind = "foot"
        else:
            big = [s for s, _ in sizes if s is not None and s >= TITLE_PT]
            is_title_ph = False
            try:
                is_title_ph = sh.is_placeholder and sh.placeholder_format.type in (1, 3)  # TITLE / CENTER_TITLE
            except Exception:
                pass
            kind = "title" if (is_title_ph or (big and top < slide_h * 0.4)) else "body"
        items.append((kind, text, sizes, top))
    return items


# --------------------------------------------------------------------------- 動き

def motion_budget(sld):
    """スライドの XML（p:sld）から動きを数える。motion.py も使う。
    returns dict(clicks, auto, per_click=[図形数,...], loops, moves)"""
    res = {"clicks": 0, "auto": 0, "per_click": [], "loops": 0, "moves": 0}
    timing = sld.find("p:timing", NS)
    if timing is None:
        return res
    main = timing.find(".//p:cTn[@nodeType='mainSeq']", NS)
    if main is None:
        return res
    lst = main.find("p:childTnLst", NS)
    if lst is None:
        return res
    for par in lst.findall("p:par", NS):
        outer = par.find("p:cTn", NS)
        auto = outer is not None and outer.find(".//p:stCondLst/p:cond[@evt='onBegin']", NS) is not None
        effs = [c for c in par.iter("{%s}cTn" % P_NS) if c.get("presetClass")]
        targets = set()
        for e in effs:
            for t in e.iter("{%s}spTgt" % P_NS):
                targets.add(t.get("spid"))
            if e.get("repeatCount") == "indefinite":
                res["loops"] += 1
            if e.get("presetClass") in ("emph", "path"):
                res["moves"] += 1
        if auto:
            res["auto"] += 1
        else:
            res["clicks"] += 1
            res["per_click"].append(len(targets))
    return res


def budget_issues(b, busy=False):
    """[(level, code, msg)]。busy なら A1〜A3 の停止を警告に下げる。"""
    out = []
    stop = "WARNING" if busy else "ERROR"
    if b["clicks"] >= BUDGET["clicks_stop"]:
        out.append((stop, "A1", f"クリックが {b['clicks']} 回（既定は2回まで・{BUDGET['clicks_stop']}回以上で停止）"))
    elif b["clicks"] >= BUDGET["clicks_warn"]:
        out.append(("WARNING", "A1", f"クリックが {b['clicks']} 回（既定は2回まで）"))
    for k, n in enumerate(b["per_click"], 1):
        if n >= BUDGET["targets_stop"]:
            out.append((stop, "A2", f"クリック{k}回目で {n} 個が動く（既定は2つまで・{BUDGET['targets_stop']}つ以上で停止）"))
        elif n >= BUDGET["targets_warn"]:
            out.append(("WARNING", "A2", f"クリック{k}回目で {n} 個が動く（既定は2つまで）"))
    if b["loops"] >= BUDGET["loops_stop"]:
        out.append((stop, "A3", f"自動で繰り返す動きが {b['loops']} つ（1枚1つまで）"))
    elif b["loops"] >= BUDGET["loops_warn"]:
        out.append(("WARNING", "A3", "自動で繰り返す動きがある（その枚の伝えたいこと自体を見せる動きか。飾りなら外す）"))
    if b["moves"] >= BUDGET["moves_warn"]:
        out.append(("WARNING", "A4", f"強調（拡大/収縮・回転）と軌跡が {b['moves']} つ（説明に要る1〜2つに絞る）"))
    return out


# --------------------------------------------------------------------------- 1枚を調べる

def check_slide(slide, no, slide_h_cm):
    issues = []
    items = text_items(slide, slide_h_cm)
    titles = [t for k, t, _, _ in items if k == "title"]
    body = [(t, s) for k, t, s, _ in items if k == "body"]
    title = " ".join(titles)
    notes = ""
    try:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
    except Exception:
        pass
    skip_text = (not titles and not body) or bool(REF_TITLE_RE.match(title))
    stat = {"slide": no, "title": title, "blocks": 0, "ja": 0, "en": 0, "sentences": 0}
    if not skip_text:
        blocks = len(body) + (1 if titles else 0)
        alltext = " ".join([title] + [t for t, _ in body])
        ja, en = measure(alltext)
        jmode = ja > en
        sentences = sum(1 for t, _ in body
                        if (measure(t)[0] >= SENTENCE_JA if is_ja(t) else measure(t)[1] >= SENTENCE_EN))
        stat.update(blocks=blocks, ja=ja, en=en, sentences=sentences)
        # I1
        if blocks >= BLOCKS_STOP:
            issues.append(("ERROR", "I1", f"読む場所が {blocks} か所（{BLOCKS_STOP}か所以上で停止・目安は8か所まで）"))
        elif blocks >= BLOCKS_WARN:
            issues.append(("WARNING", "I1", f"読む場所が {blocks} か所（目安は8か所まで）"))
        # I2
        if jmode:
            if ja > JA_STOP:
                issues.append(("ERROR", "I2", f"文字が {ja} 字（{JA_STOP}字超で停止・目安{JA_WARN}字）"))
            elif ja > JA_WARN:
                issues.append(("WARNING", "I2", f"文字が {ja} 字（目安{JA_WARN}字）"))
        else:
            if en > EN_STOP:
                issues.append(("ERROR", "I2", f"英語が {en} 語（{EN_STOP}語超で停止・目安{EN_WARN}語）"))
            elif en > EN_WARN:
                issues.append(("WARNING", "I2", f"英語が {en} 語（目安{EN_WARN}語）"))
        # I3
        tiny, small = [], 0
        for t, sizes in body:
            for pt, run in sizes:
                if pt is None:
                    continue
                if pt < FLOOR_PT:
                    tiny.append((pt, run))
                elif pt < SMALL_PT:
                    j, e = measure(run)
                    small += j if j > e else e * (SMALL_JA / SMALL_EN)
        if tiny:
            issues.append(("ERROR", "I3", f"本文に {min(p for p, _ in tiny):.1f}pt の文字（{FLOOR_PT}pt＝22px 未満）：「{tiny[0][1][:20]}」"))
        if small > SMALL_JA:
            issues.append(("WARNING", "I3", f"21pt（28px）未満の本文が多い（{'約%d語' % (small * SMALL_EN / SMALL_JA) if not jmode else '%d字' % small}）"))
        # O1
        if titles and (JOIN_EN.search(title) or JOIN_JA.search(title)):
            issues.append(("WARNING", "O1", f"見出しが2つの主張をつないでいる疑い：「{title[:60]}」→ 1つに絞るか2枚に分ける"))
        # O2
        if sentences >= SENTENCES_WARN:
            issues.append(("WARNING", "O2", f"本文に文が {sentences} 本＝伝えたいことが複数ある疑い（見出しを支える1〜2本に絞り、残りはノートへ）"))
    b = motion_budget(slide._element)
    stat.update({k: b[k] for k in ("clicks", "auto", "loops", "moves")}, per_click=b["per_click"])
    issues += budget_issues(b, busy=BUSY_MARK in notes)
    return stat, issues


def main(argv=None):
    ap = argparse.ArgumentParser(description="pptx の情報量・1枚1メッセージの手がかり・動きの予算を調べる")
    ap.add_argument("pptx")
    ap.add_argument("--slides", help="調べる枚（例 6,7）。省略で全部")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--draft", action="store_true", help="ERROR を WARNING に下げる（途中の確認用）")
    a = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    from pptx import Presentation
    prs = Presentation(a.pptx)
    want = {int(x) for x in a.slides.split(",")} if a.slides else None
    h = prs.slide_height / 360000
    rows, n_err, n_warn = [], 0, 0
    for i, s in enumerate(prs.slides, 1):
        if want and i not in want:
            continue
        stat, issues = check_slide(s, i, h)
        if a.draft:
            issues = [("WARNING" if lv == "ERROR" else lv, c, m) for lv, c, m in issues]
        n_err += sum(1 for lv, _, _ in issues if lv == "ERROR")
        n_warn += sum(1 for lv, _, _ in issues if lv == "WARNING")
        rows.append({"stat": stat, "issues": [{"level": lv, "code": c, "msg": m} for lv, c, m in issues]})
    if a.json:
        print(json.dumps({"slides": rows, "errors": n_err, "warnings": n_warn}, ensure_ascii=False, indent=1))
    else:
        for r in rows:
            st = r["stat"]
            amount = f"{st['ja']}字" if st["ja"] > st["en"] else f"{st['en']}語"
            print(f"=== {st['slide']}枚目  読む場所{st['blocks']}・{amount}・文{st['sentences']}本"
                  f"・クリック{st['clicks']}{'＋自動' if st['auto'] else ''}{st['per_click'] if st['per_click'] else ''}"
                  f"  {st['title'][:50]}")
            for it in r["issues"]:
                print(f"  {it['level']:7} {it['code']} {it['msg']}")
        print(f"ERROR {n_err} / WARNING {n_warn}")
    return 1 if n_err else 0


if __name__ == "__main__":
    raise SystemExit(main())
