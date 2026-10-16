"""実験レポート用 rich-text ヘルパ（report スキル STEP 2d の実体）。

DSL: `V_FB` / `V_{p-p}` を下付き、`x^{2}` / `10^{-3}` を上付きとして
python-docx の run に展開する。プレーンな `paragraph.add_run(text)` を
直接呼ぶと下付き・上付きが失われるため、本文・式・表ヘッダ・キャプションは
必ずこのモジュール経由で書く。

使い方:
    import sys, os; sys.path.insert(0, os.path.expanduser(r'~/.claude/skills/report/scripts'))
    from rich_text import add_rich, set_jp_font
    p = doc.add_paragraph()
    add_rich(p, "ツェナー電圧 V_Z は 10^{-3} A で測定した。")

生成後検証: docx を開き直し、物理量段落に `run.font.subscript == True` の
run が1つ以上あることを確認する（0件ならプレーン残存 = NG）。
"""
import re

from docx.oxml.ns import qn
from docx.shared import Pt

# 下付き: 英字/ギリシャ文字1字 + `_` + {複合} or 英数字連続
#   ただし base の直前が英数字・`_` のとき（read_adc / SPI_MOSI / RPi_CAE のような識別子）は下付きにしない
#   （2026-09-29 テーマ9 の本文で read_adc() が「read＋下付き adc」になった事故の対策）
# 上付き: `^{...}`
_TOKEN_PATTERN = re.compile(
    r"(?<![A-Za-z0-9_])(?P<base>[A-Za-zβαρΩμφω]̇?)_(?P<sub>\{[^}]+\}|[A-Za-z0-9]+)"
    r"|\^\{(?P<sup>[^}]+)\}"
)


def set_jp_font(run, font_jp: str = "游明朝") -> None:
    """run の日本語(eastAsia)フォントを設定する。"""
    run.font.name = font_jp
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_jp)


def _plain_run(paragraph, text, font_size, font_jp, bold):
    r = paragraph.add_run(text)
    r.font.size = Pt(font_size)
    r.bold = bold
    set_jp_font(r, font_jp)
    return r


def add_rich(paragraph, text: str, font_size: float = 10.5,
             font_jp: str = "游明朝", bold: bool = False) -> None:
    """`X_Y` / `X_{YZ}` を下付き、`^{Y}` を上付きとして描画する。

    下付き・上付きを出現順に一括処理する（旧実装は最後の下付きより前に
    ある上付きを取りこぼしていた）。
    """
    pos = 0
    for m in _TOKEN_PATTERN.finditer(text):
        if m.start() > pos:
            _plain_run(paragraph, text[pos:m.start()], font_size, font_jp, bold)
        if m.group("base") is not None:
            _plain_run(paragraph, m.group("base"), font_size, font_jp, bold)
            sub = m.group("sub")
            if sub.startswith("{"):
                sub = sub[1:-1]
            r = _plain_run(paragraph, sub, font_size, font_jp, bold)
            r.font.subscript = True
        else:
            r = _plain_run(paragraph, m.group("sup"), font_size, font_jp, bold)
            r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        _plain_run(paragraph, text[pos:], font_size, font_jp, bold)


# 旧ドキュメントの名前でも呼べるようにする
_add_rich = add_rich
