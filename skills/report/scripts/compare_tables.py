"""report/compare_tables.py - 入力（実測データ）と出力レポートの表を数値セル単位で突き合わせる

usage:
    py compare_tables.py <出力.docx> --source <入力.xlsx|入力.docx> [--source ...]
                         [--skip-table N ...] [--skip-col N:k ...]

目的:
    SKILL.md「データ改ざん禁止」の機械検査。出力 docx の表に入っている数値が、
    入力ファイル（実測を記録した xlsx / docx）にそのまま在るかを1セルずつ確かめる。
    出力にあって入力に無い数値・値の食い違いを列挙し、1件でもあれば exit 1。

判定（1つの数値ごと）:
    [NG]   入力のどこにも無い数値（捏造・転記ミスの疑い）
    [NG]   入力が手入力の値なのに、丸めて桁数を変えている（8.32 → 8.3）
    [NG]   列の対応がはっきり取れた列で、同じ位置の入力値と違う（値の食い違い・行の抜き取り）
    [NG]   列の対応がはっきり取れた列で、入力では空欄のセルに数値が入っている
    [警告] 値は同じだが末尾の 0 の桁数が違う（1.3 と 1.30。Excel の General は末尾0を消すため）
    [警告] 単位の換算（×1000 等）でだけ一致する
    [警告] 見出し的なセル（「IB＝5 μA」「R = 1 kΩ」など文字の混じるセル）の数値が入力に無い
    [情報] Excel の数式セル（計算列）の値を丸めて一致した

数値の扱い（誤判定を減らすための正規化）:
    - 全角→半角、マイナス記号の揺れ（− ‐ – 等）を統一。「0-10」のように数字の直後の - は区切り
    - 1,000 の桁区切りを外す。1.2e-3・1.2×10^-3（Word の上付き書式・Unicode ⁻³ も）を1つの数値に
    - R1・CH2・2SC1815 のように英字の直後の数字は記号の一部として数えない。下付きの添字も数えない
    - 比較は Decimal（2進小数の誤差なし）。桁数は「小数点以下の桁数」まで見る
    - xlsx は「保存されている値」と「表示書式（0.00 等）で見える値」の両方を候補にする

検査しない表（自動）:
    - 使用機器の表（キャプションか1行目に「機器」）／プログラムのリスト（「リスト」・#include・; の多い表）
    - --skip-table N で指定した表N（理論値の表など、入力に無いのが正しい表）
    - --skip-col N:k で指定した 表N の k 列目（1始まり。誤差率などの計算列）
    除外した表・列は必ず結果に一覧表示する（黙って外さない）。

exit: 0=NGなし / 1=NGあり / 2=使い方・ファイルの誤り
"""
import sys
import io
import os
import re
import bisect
import argparse
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, ROUND_HALF_EVEN

from docx import Document
from docx.oxml.ns import qn

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass


# ============================================================
# 数値の取り出し
# ============================================================

# Unicode の上付き文字 → 普通の文字（10⁻³ を 10^(-3) に直すため）
SUP_MAP = {
    '⁰': '0', '¹': '1', '²': '2', '³': '3', '⁴': '4',
    '⁵': '5', '⁶': '6', '⁷': '7', '⁸': '8', '⁹': '9',
    '⁺': '+', '⁻': '-',
}
# Unicode の下付き文字（添字なので数値として数えない）
SUB_CHARS = '₀₁₂₃₄₅₆₇₈₉₊₋'
# マイナス記号の揺れ（em dash「—」はデータ無しの印に使うので含めない）
MINUS_CHARS = ['−', '‐', '‑', '‒', '–', '﹣']

# 数値1個の形。直前が英数字・小数点・下線なら記号の一部（R1, 2SC1815, No.1）なので数えない
NUM_RE = re.compile(
    r'(?<![A-Za-z0-9._])'
    r'(?P<sign>-)?'
    r'(?P<mant>\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)'
    r'(?:(?P<e>E[-+]?\d+)(?![A-Za-z])'
    r'|\s*[×xX*]\s*10\s*\^\s*\(\s*(?P<p10>[-+]?\d+)\s*\)'
    r'|\^\s*\(\s*(?P<pow>[-+]?\d+)\s*\))?'
)

# 数値と単位だけのセルかを見るための単位の一覧（これ以外の文字が残れば「見出し的なセル」）
UNIT_WORDS = ['kHz', 'MHz', 'Hz', 'mV', 'μV', 'kV', 'V', 'mA', 'μA', 'nA', 'A',
              'kΩ', 'MΩ', 'Ω', 'μF', 'nF', 'pF', 'F', 'mH', 'μH', 'H', 'ms', 'μs', 'ns',
              'mW', 'W', 'dB', 's', '°', '%', 'rad', 'K', '℃']


class Num(object):
    """表の中の数値1個。value は Decimal、text は元の書き方"""

    def __init__(self, value, text):
        self.value = value
        self.text = text

    def exponent(self):
        # 小数点以下の桁数の目安（0.100 → -3、8.3E-3 → -4、100 → 0）
        return self.value.as_tuple().exponent


def normalize_text(text):
    """数値を拾いやすい形にそろえる（上付き・下付き・全角・マイナスの揺れ）"""
    out = []
    i = 0
    while i < len(text):
        ch = text[i]
        if ch in SUP_MAP:
            # 上付きが続く所をまとめて ^(...) にする
            buf = ''
            while i < len(text) and text[i] in SUP_MAP:
                buf += SUP_MAP[text[i]]
                i += 1
            out.append('^(' + buf + ')')
            continue
        if ch in SUB_CHARS:
            i += 1
            continue
        out.append(ch)
        i += 1
    text = ''.join(out)
    text = unicodedata.normalize('NFKC', text)
    for m in MINUS_CHARS:
        text = text.replace(m, '-')
    # 指数の e を大文字にそろえる（1.2e-3 → 1.2E-3）
    text = re.sub(r'(?<=\d)e(?=[-+]?\d)', 'E', text)
    # 数字・英字の直後の - は区切り（0-10 は「0 から 10」。E の直後は指数なので除く）
    text = re.sub(r'(?<=[0-9A-DF-Za-z])-(?=\d)', ' ', text)
    # 下付き書式の添字 _(...) は数えない
    text = re.sub(r'_\([^)]*\)', ' ', text)
    # 10 に付いていない上付き（単位の m^(2) など）は数えない
    text = re.sub(r'(?<!10)\^\([^)]*\)', ' ', text)
    return text


def parse_numbers(text):
    """文字列から数値を全部拾って Num のリストで返す"""
    text = normalize_text(text)
    nums = []
    for m in NUM_RE.finditer(text):
        mant = m.group('mant').replace(',', '')
        try:
            if m.group('e'):
                value = Decimal(mant + m.group('e'))
            elif m.group('p10'):
                value = Decimal(mant + 'E' + m.group('p10'))
            elif m.group('pow') and mant == '10':
                value = Decimal('1E' + m.group('pow'))
            else:
                value = Decimal(mant)
        except InvalidOperation:
            continue
        if m.group('sign'):
            value = -value
        nums.append(Num(value, m.group(0).strip()))
    return nums


def is_label_cell(text):
    """数値と単位以外の文字が残るセル（「IB＝5 μA」「R = 1 kΩ」）＝見出し的なセル"""
    t = normalize_text(text)
    t = NUM_RE.sub(' ', t)
    for u in sorted(UNIT_WORDS, key=len, reverse=True):
        t = t.replace(u, ' ')
    t = re.sub(r'[\s\[\]()（）\-+.,:;=×/～〜~±]', '', t)
    return len(t) > 0


# ============================================================
# Word の表を読む
# ============================================================

def run_aware_text(el):
    """段落・セルの文字を、上付き書式は ^(..)、下付き書式は _(..) の印を付けて返す"""
    lines = []
    for p in el.iter(qn('w:p')):
        parts = []
        for r in p.iter(qn('w:r')):
            t = ''.join(x.text or '' for x in r.iter(qn('w:t')))
            if not t:
                continue
            va = None
            rpr = r.find(qn('w:rPr'))
            if rpr is not None:
                v = rpr.find(qn('w:vertAlign'))
                if v is not None:
                    va = v.get(qn('w:val'))
            if va == 'superscript':
                parts.append('^(' + t + ')')
            elif va == 'subscript':
                parts.append('_(' + t + ')')
            else:
                parts.append(t)
        lines.append(''.join(parts))
    return '\n'.join(lines)


def plain_text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def table_grid(tbl):
    """w:tbl を2次元の表にする。横の結合は広げ、縦結合の続きのセルは None（値を二重に数えない）。
    戻り値: (grid, header_marks)  grid[r][c] = (表示用の文字, 数値拾い用の文字) or None"""
    grid = []
    header_marks = []
    for tr in tbl.findall(qn('w:tr')):
        row = []
        is_header = False
        trpr = tr.find(qn('w:trPr'))
        if trpr is not None and trpr.find(qn('w:tblHeader')) is not None:
            is_header = True
        for tc in tr.findall(qn('w:tc')):
            span = 1
            vcont = False
            tcpr = tc.find(qn('w:tcPr'))
            if tcpr is not None:
                gs = tcpr.find(qn('w:gridSpan'))
                if gs is not None:
                    span = int(gs.get(qn('w:val')))
                vm = tcpr.find(qn('w:vMerge'))
                if vm is not None and vm.get(qn('w:val')) != 'restart':
                    vcont = True
            if vcont:
                cell = None
            else:
                cell = (plain_text(tc), run_aware_text(tc))
            row.append(cell)
            for _ in range(span - 1):
                row.append(None)
        grid.append(row)
        header_marks.append(is_header)
    return grid, header_marks


def docx_tables_with_captions(doc):
    """本文直下の表を順に、直前のキャプション（表N／リストN）と一緒に返す"""
    out = []
    last_para = ''
    for child in list(doc.element.body):
        if child.tag == qn('w:p'):
            text = plain_text(child).strip()
            if text:
                last_para = text
        elif child.tag == qn('w:tbl'):
            caption = ''
            if re.match(r'^(表|リスト)\s?\d+', last_para):
                caption = last_para
            out.append((caption, child))
            last_para = ''
    return out


# ============================================================
# 入力（実測データ）を読む
# ============================================================

class SrcCell(object):
    """入力の1セル。cands は比べる候補の Decimal（保存値と表示値）"""

    def __init__(self, where, cands, is_formula, text):
        self.where = where
        self.cands = cands
        self.is_formula = is_formula
        self.text = text
        # Excel の General 書式の数値は末尾の0を保存できない（1.30 と打っても 1.3 になる）。
        # そういうセルは 1.3 と 1.30 の違いを桁数の違いとみなさない
        self.trailing_ok = False
        # 表の中の位置（grid番号・行・列）。本文段落の数値は位置なし＝None
        self.gi = None
        self.r = None
        self.c = None


def excel_display(value, fmt):
    """Excel の表示書式のうちよくある形（0 / 0.00 / 0.0% / 0.00E+00）だけ真似る。
    分からない書式は None（保存値だけで比べる）"""
    if not fmt or fmt == 'General':
        return None
    f = fmt.split(';')[0]
    f = re.sub(r'"[^"]*"', '', f)
    f = re.sub(r'\[[^\]]*\]', '', f)
    f = f.replace('\\', '')
    if '#' in f.split('.')[-1] and '.' in f:
        return None
    m = re.search(r'0(?:\.(0+))?', f)
    if not m:
        return None
    decimals = len(m.group(1)) if m.group(1) else 0
    try:
        d = Decimal(repr(value))
    except InvalidOperation:
        return None
    if 'E+' in f.upper() or 'E-' in f.upper():
        s = ('{:.' + str(decimals) + 'E}').format(d)
        return Decimal(s)
    if '%' in f:
        d = d * 100
    q = Decimal(1).scaleb(-decimals)
    return d.quantize(q, rounding=ROUND_HALF_UP)


def read_xlsx(path):
    """xlsx の全シートを読む。戻り値: (grids, cells, 数式でキャッシュ値が無いセル数)"""
    import openpyxl
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    name = os.path.basename(path)
    grids = []
    cells = []
    no_cache = 0
    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        grid = []
        for row in ws.iter_rows():
            grow = []
            for c in row:
                f = wf.cell(row=c.row, column=c.column).value
                is_formula = isinstance(f, str) and f.startswith('=')
                v = c.value
                where = '%s %s!%s' % (name, ws.title, c.coordinate)
                if v is None:
                    if is_formula:
                        no_cache += 1
                    grow.append(None)
                    continue
                if isinstance(v, bool):
                    grow.append(None)
                    continue
                if isinstance(v, (int, float)):
                    cands = [Decimal(repr(v))]
                    disp = excel_display(v, c.number_format)
                    if disp is not None:
                        cands.append(disp)
                    sc = SrcCell(where, cands, is_formula, str(v))
                    if disp is None:
                        sc.trailing_ok = True
                elif isinstance(v, str):
                    nums = parse_numbers(v)
                    if not nums:
                        grow.append(SrcCell(where, [], False, v))
                        continue
                    sc = SrcCell(where, [n.value for n in nums], False, v)
                else:
                    # 日付など
                    grow.append(None)
                    continue
                grow.append(sc)
                cells.append(sc)
            grid.append(grow)
        grids.append(('%s %s' % (name, ws.title), grid))
    return grids, cells, no_cache


def read_docx_source(path):
    """入力が docx（プレレポート・記録表）のとき。表は位置付きで、段落の数値は位置なしで読む"""
    doc = Document(path)
    name = os.path.basename(path)
    grids = []
    cells = []
    k = 0
    for caption, tbl in docx_tables_with_captions(doc):
        k += 1
        label = caption.split()[0] if caption else '表#%d' % k
        grid_raw, _marks = table_grid(tbl)
        grid = []
        for r, row in enumerate(grid_raw):
            grow = []
            for c, cell in enumerate(row):
                if cell is None:
                    grow.append(None)
                    continue
                nums = parse_numbers(cell[1])
                where = '%s %s(%d行%d列)' % (name, label, r + 1, c + 1)
                sc = SrcCell(where, [n.value for n in nums], False, cell[0])
                grow.append(sc)
                if nums:
                    cells.append(sc)
            grid.append(grow)
        grids.append(('%s %s' % (name, label), grid))
    # 本文段落の数値（位置合わせには使わず、在るかどうかだけに使う）
    for i, p in enumerate(doc.element.body.findall(qn('w:p'))):
        nums = parse_numbers(run_aware_text(p))
        if nums:
            cells.append(SrcCell('%s 本文段落%d' % (name, i + 1), [n.value for n in nums], False, ''))
    return grids, cells


# ============================================================
# 一致の判定
# ============================================================

EXACT = 'EXACT'        # 値も桁数も同じ
SAME = 'SAME'          # 値は同じで末尾の0の桁数だけ違う
ROUND_F = 'ROUND_F'    # 数式セルの値を丸めると一致
ROUND_C = 'ROUND_C'    # 手入力の値を丸めると一致（桁数を変えている）
SCALED = 'SCALED'      # 単位換算（×10^±3, ±6）で一致
NONE = 'NONE'          # どこにも無い


def match_one(out_value, src_value, is_formula):
    """出力の数値1個と入力の候補1個を比べて、一致の種類を返す（一致しなければ None）"""
    if out_value == src_value:
        if out_value.as_tuple().exponent == src_value.as_tuple().exponent:
            return EXACT
        return SAME
    e_out = out_value.as_tuple().exponent
    e_src = src_value.as_tuple().exponent
    if isinstance(e_out, int) and isinstance(e_src, int) and e_src < e_out:
        q = Decimal(1).scaleb(e_out)
        for mode in (ROUND_HALF_UP, ROUND_HALF_EVEN):
            if src_value.quantize(q, rounding=mode) == out_value:
                if is_formula:
                    return ROUND_F
                return ROUND_C
    return None


def match_cell(out_value, src_value, sc):
    """match_one に入力セルの事情を足したもの。
    0 は 0.00 と書いても 0 なので桁数を問わない。General 書式の Excel 数値も末尾の0は問わない"""
    kind = match_one(out_value, src_value, sc.is_formula)
    if kind == SAME:
        if out_value == 0 or sc.trailing_ok:
            return EXACT
    return kind


RANK = {EXACT: 0, SAME: 1, ROUND_F: 2, SCALED: 3, ROUND_C: 4, NONE: 5}


class SourceIndex(object):
    """入力の全数値を値の大きさ順に並べ、近い値だけを調べて速く探す"""

    def __init__(self, cells):
        self.items = []
        for sc in cells:
            for v in sc.cands:
                self.items.append((float(v), v, sc))
        self.items.sort(key=lambda x: x[0])
        self.keys = [x[0] for x in self.items]

    def near(self, value, width):
        lo = bisect.bisect_left(self.keys, float(value) - width)
        hi = bisect.bisect_right(self.keys, float(value) + width)
        return self.items[lo:hi]

    def best_match(self, num):
        """出力の数値 num に一番よく合う入力を探す。戻り値: (種類, 入力セル or None)"""
        best = (NONE, None)
        e = num.value.as_tuple().exponent
        if not isinstance(e, int):
            return best
        width = float(Decimal(1).scaleb(e)) * 0.51 + 1e-12
        for _f, v, sc in self.near(num.value, width):
            kind = match_cell(num.value, v, sc)
            if kind is not None and RANK[kind] < RANK[best[0]]:
                best = (kind, sc)
                if kind == EXACT:
                    return best
        if best[0] != NONE:
            return best
        # 単位換算（A と mA など）でだけ合うか
        for k in (3, -3, 6, -6):
            scaled = num.value.scaleb(-k)
            w = float(Decimal(1).scaleb(e - k)) * 0.51 + 1e-15
            for _f, v, sc in self.near(scaled, w):
                kind = match_cell(scaled, v, sc)
                if kind in (EXACT, SAME, ROUND_F):
                    return (SCALED, sc)
        return best


# ============================================================
# 列の位置合わせ（値の食い違い・空欄への記入を見つける）
# ============================================================

def align_column(col_cells, index, prefer=None):
    """出力の1列（[(行, Num)]）が入力のどの列（または行）に対応するかを多数決で決める。
    出力の値と一致する入力セルごとに「入力の列○で行が△ずれている」に1票入れ、
    一番票の多い対応を返す。表を縦横入れ替えた場合（入力の行に対応）も数える。
    票が同数のときは prefer（同じ表の他の列と同じ「表・向き・ずれ」）を優先する。
    位相差のように値の並びが別の列でもずれて現れる表で、隣の列に取り違えないため。
    戻り値: (grid番号, 'col'/'row', 列 or 行, ずれ, 票数) または None"""
    votes = {}
    for r, num in col_cells:
        e = num.value.as_tuple().exponent
        if not isinstance(e, int):
            continue
        width = float(Decimal(1).scaleb(e)) * 0.51 + 1e-12
        for _f, v, sc in index.near(num.value, width):
            if sc.gi is None:
                continue
            kind = match_cell(num.value, v, sc)
            if kind not in (EXACT, SAME, ROUND_F):
                continue
            for key in ((sc.gi, 'col', sc.c, sc.r - r), (sc.gi, 'row', sc.r, sc.c - r)):
                if key not in votes:
                    votes[key] = set()
                votes[key].add(r)
    best = None
    for key, rows in votes.items():
        if best is None or len(rows) > len(votes[best]):
            best = key
        elif len(rows) == len(votes[best]) and prefer is not None:
            if (key[0], key[1], key[3]) == prefer and (best[0], best[1], best[3]) != prefer:
                best = key
    if best is None:
        return None
    return best + (len(votes[best]),)


def src_at(grids, key, r):
    """位置合わせの結果 key で、出力の行 r に対応する入力セルを返す（範囲外は 'OUT'）"""
    gi, mode, fixed, off = key[0], key[1], key[2], key[3]
    grid = grids[gi][1]
    if mode == 'col':
        sr, sc_i = r + off, fixed
    else:
        sr, sc_i = fixed, r + off
    if sr < 0 or sr >= len(grid):
        return 'OUT'
    row = grid[sr]
    if sc_i < 0 or sc_i >= len(row):
        return 'OUT'
    return row[sc_i]


def where_of(grids, key, r):
    gi, mode, fixed, off = key[0], key[1], key[2], key[3]
    if mode == 'col':
        return '%s の %d行%d列' % (grids[gi][0], r + off + 1, fixed + 1)
    return '%s の %d行%d列' % (grids[gi][0], fixed + 1, r + off + 1)


# ============================================================
# 本体
# ============================================================

def header_row_count(grid, marks):
    """見出し行の数: tblHeader の付いた先頭行、無ければ1行目。
    さらに続く行が全部「見出し的なセル」なら見出しとみなす（最大3行）"""
    n = 0
    while n < len(marks) and marks[n]:
        n += 1
    if n == 0:
        n = 1
    while n < len(grid) and n < 3:
        cells = [c for c in grid[n] if c is not None and c[0].strip()]
        if cells and all(is_label_cell(c[0]) for c in cells):
            n += 1
        else:
            break
    return n


def looks_like_code(caption, grid):
    if caption.startswith('リスト'):
        return True
    text = '\n'.join(c[0] for row in grid for c in row if c is not None)
    if '#include' in text or text.count(';') >= 3:
        return True
    return False


def parse_skip_cols(values):
    out = {}
    for s in values or []:
        m = re.match(r'^\s*(\d+)\s*:\s*(\d+)\s*$', s)
        if not m:
            raise ValueError('--skip-col は 表番号:列番号（例 2:5）で指定: %r' % s)
        out.setdefault(int(m.group(1)), set()).add(int(m.group(2)))
    return out


def read_sources(sources):
    """入力ファイルを全部読み、(grids, cells, 警告) を返す。セルには表の中の位置を付ける"""
    grids = []
    cells = []
    warns = []
    for s in sources:
        ext = os.path.splitext(s)[1].lower()
        if ext in ('.xlsx', '.xlsm'):
            g, c, no_cache = read_xlsx(s)
            if no_cache:
                warns.append('[警告] %s: 数式セル %d 個にキャッシュ値が無い'
                             '（Excel で開いて保存し直すと計算列も照合できる）'
                             % (os.path.basename(s), no_cache))
        elif ext == '.docx':
            g, c = read_docx_source(s)
        else:
            raise ValueError('入力は .xlsx / .xlsm / .docx のみ対応: %s' % s)
        grids += g
        cells += c
    # 位置合わせ用に、各セルへ「何番目の表の何行何列か」を書き込む
    for gi, (_name, grid) in enumerate(grids):
        for r, row in enumerate(grid):
            for c, sc in enumerate(row):
                if sc is not None:
                    sc.gi = gi
                    sc.r = r
                    sc.c = c
    return grids, cells, warns


def judge_aligned(num, src):
    """位置合わせで対応した入力セル src と出力の数値 num を比べる。
    戻り値: 一致の種類（EXACT/SAME/ROUND_F/ROUND_C）か None（食い違い）"""
    best = None
    for v in src.cands:
        kind = match_cell(num.value, v, src)
        if kind is not None and (best is None or RANK[kind] < RANK[best]):
            best = kind
    return best


def judge_cell(num, loc, label_cell, index, grids, key, r):
    """出力の数値1個の判定。戻り値: ('NG'/'WARN'/'INFO'/'OK', 文)
    key は列の位置合わせの結果（はっきり対応が取れた列だけ。無ければ None）"""
    kind, sc = index.best_match(num)
    elsewhere = ''
    if kind in (EXACT, SAME) and sc is not None:
        elsewhere = '（この値は入力の別の場所 %s にある＝行の抜き取り・並べ替えの可能性）' % sc.where

    # (1) 位置合わせできた列: 同じ位置の入力セルと比べる
    if key is not None:
        src = src_at(grids, key, r)
        if src == 'OUT' or src is None:
            return ('NG', '%s: 入力の同じ位置（%s）は空欄なのに、出力には数値がある%s'
                    % (loc, where_of(grids, key, r), elsewhere))
        if not src.cands:
            return ('NG', '%s: 入力の同じ位置 %s は「%s」（数値でない）なのに、出力には数値がある%s'
                    % (loc, src.where, src.text, elsewhere))
        here = judge_aligned(num, src)
        if here is None:
            return ('NG', '%s: 値の食い違い。入力の同じ位置 %s は「%s」%s'
                    % (loc, src.where, src.text, elsewhere))
        if here == ROUND_C and kind != EXACT:
            return ('NG', '%s: 入力の同じ位置 %s の手入力値「%s」を丸めて桁数を変えている'
                    % (loc, src.where, src.text))
        # 同じ位置で一致した → その一致の種類で下の判定へ。
        # ただし入力の別の場所（書式付きの清書シート等）に桁数まで同じ値があれば、そちらを採る
        if RANK[here] <= RANK[kind]:
            kind = here
            sc = src

    # (2) 入力全体から探した結果で判定
    if kind == NONE:
        if label_cell:
            return ('WARN', '%s: 見出し的なセルの数値が入力に無い（条件値なら問題なし）' % loc)
        return ('NG', '%s: 入力のどこにも無い数値' % loc)
    if kind == ROUND_C:
        return ('NG', '%s: 入力に同じ値が無い。丸めると合うのは手入力値「%s」（%s）だけ'
                '＝桁数を変えた転記か、別の値' % (loc, sc.text, sc.where))
    if kind == SCALED:
        return ('WARN', '%s: 単位換算でだけ一致（%s「%s」）。単位の書き間違いでないか確認'
                % (loc, sc.where, sc.text))
    if kind == SAME:
        return ('WARN', '%s: 値は同じだが末尾の0の桁数が違う（%s「%s」）' % (loc, sc.where, sc.text))
    if kind == ROUND_F:
        return ('INFO', '%s: 数式セル %s の値を丸めて一致' % (loc, sc.where))
    return ('OK', '')


def compare(out_path, sources, skip_tables=None, skip_cols=None):
    """戻り値: (ng, warn, info, notes)  それぞれ文字列のリスト"""
    skip_tables = set(skip_tables or [])
    skip_cols = skip_cols or {}
    ng, warn, info, notes = [], [], [], []

    grids, cells, warns = read_sources(sources)
    warn += warns
    index = SourceIndex(cells)
    notes.append('入力: %s（数値セル %d 個）' % (', '.join(os.path.basename(s) for s in sources), len(cells)))

    doc = Document(out_path)
    k = 0
    for caption, tbl in docx_tables_with_captions(doc):
        k += 1
        m = re.match(r'^表\s?(\d+)', caption)
        tnum = int(m.group(1)) if m else None
        if tnum is not None:
            label = '表%d' % tnum
        else:
            label = '（キャプション無しの表 #%d）' % k
        grid, marks = table_grid(tbl)
        first_row = ' '.join(c[0] for c in (grid[0] if grid else []) if c is not None)
        if '機器' in caption or '機器' in first_row:
            notes.append('除外: %s（使用機器の表）' % label)
            continue
        if looks_like_code(caption, grid):
            notes.append('除外: %s（プログラムのリスト）' % label)
            continue
        if tnum is not None and tnum in skip_tables:
            notes.append('除外: %s（--skip-table で指定）' % label)
            continue
        cols_skip = set()
        if tnum is not None:
            cols_skip = skip_cols.get(tnum, set())
        for c in sorted(cols_skip):
            notes.append('除外: %s の %d列目（--skip-col で指定）' % (label, c))

        # 見出し行より下のセルから数値を拾う
        nh = header_row_count(grid, marks)
        found = []          # (行, 列, Num, 見出し的なセルか, そのセルの数値の個数)
        columns = {}        # 列 → [(行, Num)]  数値1個だけの普通のセル（位置合わせ用）
        for r in range(nh, len(grid)):
            for c, cell in enumerate(grid[r]):
                if cell is None or (c + 1) in cols_skip:
                    continue
                nums = parse_numbers(cell[1])
                if not nums:
                    continue
                label_cell = is_label_cell(cell[0])
                for n in nums:
                    found.append((r, c, n, label_cell, len(nums)))
                if len(nums) == 1 and not label_cell:
                    columns.setdefault(c, []).append((r, nums[0]))

        # 列の位置合わせ。票数3以上かつ列の数値の8割以上が同じ位置で一致した列だけ
        # 「はっきり対応が取れた列」として、同じ位置どうしで比べる
        # 1回目で表全体として一番多い「表・向き・ずれ」を求め、2回目で同数票の時にそれを優先する
        prefer = None
        for round_no in (1, 2):
            aligned = {}
            for c, col_cells in columns.items():
                res = align_column(col_cells, index, prefer)
                if res is None:
                    continue
                nvote = res[4]
                if nvote >= 3 and nvote >= 0.8 * len(col_cells):
                    aligned[c] = res[:4]
            count = {}
            for key in aligned.values():
                t = (key[0], key[1], key[3])
                count[t] = count.get(t, 0) + 1
            for t in count:
                if prefer is None or count[t] > count[prefer]:
                    prefer = t

        n_ok = 0
        for r, c, num, label_cell, count in found:
            loc = '%s %d行%d列「%s」' % (label, r + 1, c + 1, num.text)
            key = None
            if c in aligned and not label_cell and count == 1:
                key = aligned[c]
            level, msg = judge_cell(num, loc, label_cell, index, grids, key, r)
            if level == 'NG':
                ng.append('[NG] ' + msg)
            elif level == 'WARN':
                warn.append('[警告] ' + msg)
            elif level == 'INFO':
                info.append('[情報] ' + msg)
            else:
                n_ok += 1

        align_desc = []
        for c in sorted(aligned):
            key = aligned[c]
            if key[1] == 'col':
                align_desc.append('%d列→%s の %d列' % (c + 1, grids[key[0]][0], key[2] + 1))
            else:
                align_desc.append('%d列→%s の %d行' % (c + 1, grids[key[0]][0], key[2] + 1))
        tail = ''
        if align_desc:
            tail = '／列の対応: ' + '、'.join(align_desc)
        notes.append('検査: %s 数値 %d 個（完全一致 %d）%s' % (label, len(found), n_ok, tail))
    return ng, warn, info, notes


def print_list(title, items, limit):
    if not items:
        return
    print('%s（%d件）' % (title, len(items)))
    for s in items[:limit]:
        print('  ' + s)
    if len(items) > limit:
        print('  ... 他 %d 件' % (len(items) - limit))
    print()


def main():
    ap = argparse.ArgumentParser(description='入力（実測）と出力 docx の表を数値セル単位で突き合わせる')
    ap.add_argument('docx', help='出力レポート（.docx）')
    ap.add_argument('--source', action='append', required=True,
                    help='入力ファイル（.xlsx/.xlsm/.docx）。複数指定可')
    ap.add_argument('--skip-table', action='append', type=int, default=[],
                    help='検査しない表の番号（理論値の表など）。複数指定可')
    ap.add_argument('--skip-col', action='append', default=[],
                    help='検査しない列（表番号:列番号、1始まり。例 2:5）。複数指定可')
    ap.add_argument('--max-lines', type=int, default=40, help='種類ごとの表示件数の上限')
    args = ap.parse_args()

    for p in [args.docx] + args.source:
        if not os.path.exists(p):
            print('[NG] ファイルが見つかりません: %s' % p)
            sys.exit(2)
    try:
        skip_cols = parse_skip_cols(args.skip_col)
        ng, warn, info, notes = compare(args.docx, args.source, args.skip_table, skip_cols)
    except ValueError as e:
        print('[NG] %s' % e)
        sys.exit(2)

    print('表データ照合: %s\n' % args.docx)
    for s in notes:
        print('  ' + s)
    print()
    print_list('[NG] 入力と一致しない数値', ng, args.max_lines)
    print_list('[警告] 確認が要る数値', warn, args.max_lines)
    print_list('[情報] 計算列の丸め', info, 10)
    if ng:
        print('[NG] %d 件。出力の表を入力の値に戻す（捏造・改ざん禁止 P-035）。' % len(ng))
        print('     理論値の表など入力に無いのが正しい表は --skip-table N、計算列は --skip-col N:k で除外し、')
        print('     除外した理由を報告に書く。')
        sys.exit(1)
    print('[OK] 出力の表の数値はすべて入力にある（NG 0 件・警告 %d 件）' % len(warn))
    sys.exit(0)


if __name__ == '__main__':
    main()
