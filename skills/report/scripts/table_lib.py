"""実験レポート用の Word 表・図挿入ヘルパ（report スキル STEP 3 の実体）。

表の書式（キャプション上・黒罫線・マージヘッダ・下付き文字）は毎回
XML 手書きで再実装されてきたため、python-docx ベースでライブラリ化した。
罫線（外枠 sz=8・内側 sz=4 の黒単線）・表全体の中央配置・fixed レイアウトは
過去テーマ5の表4と同方式。セル文字は全て rich_text.add_rich 経由なので
`V_{out}` の下付きや `10^{-3}` の上付きがそのまま使える。

使い方:
    import sys, os; sys.path.insert(0, os.path.expanduser(r'~/.claude/skills/report/scripts'))
    from table_lib import add_table_caption, build_table, add_figure

    # 表（キャプションは表の上）
    add_table_caption(doc, "表1　使用機器")
    build_table(doc,
                header_rows=[["機器・部品", "仕様"]],
                data_rows=[["直流電源", "PMC18-2"],
                           ["抵抗", "1 kΩ、10 kΩ"]],
                col_widths_cm=[5.0, 7.0])

    # 2段マージヘッダ（dict で colspan / rowspan を指定）
    add_table_caption(doc, "表2　周波数特性の測定記録")
    build_table(doc,
                header_rows=[
                    [{"text": "周波数 f [Hz]", "rowspan": 2},
                     {"text": "出力電圧 V_{out} [V]", "colspan": 3}],
                    ["R=1 kΩ", "R=10 kΩ", "R=100 kΩ"],
                ],
                data_rows=[["1000", "", "", ""]])

    # 図（キャプションは図の下）
    add_figure(doc, "figures/sin.png", "図1　正弦波の例", width_cm=13)

注意: 実測データが無いセルは空文字 "" のまま渡す。空セルは空のまま
出力される（理論値などで勝手に埋めない。データ捏造禁止ルールの根幹）。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from rich_text import add_rich

# レポート本文と同じ明朝体（キャプション・表セル共通）
FONT_JP = "ＭＳ 明朝"


def add_table_caption(doc, text, font_size=10.5):
    """「表N　キャプション」段落を追加する。表を作る前に呼ぶ（表の上）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True   # キャプションと表を別ページに離さない（2026-09-29 テーマ1 表8で離れた）
    add_rich(p, text, font_size=font_size, font_jp=FONT_JP)
    return p


def add_figure(doc, png_path, caption, width_cm=13):
    """図を中央配置で挿入し、キャプション「図N　…」を図の下に置く。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(png_path, width=Cm(width_cm))
    p.paragraph_format.keep_with_next = True   # 図とキャプションを別ページに離さない（2026-09-29 テーマ5 図5で離れた）
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    add_rich(cap, caption, font_size=10.5, font_jp=FONT_JP)
    return cap


def build_table(doc, header_rows, data_rows, col_widths_cm=None,
                repeat_header=True, font_size=10):
    """罫線・中央配置つきの表を作る。

    header_rows: ヘッダ行のリスト。各セルは文字列か
                 {"text": ..., "colspan": n, "rowspan": n} の dict。
    data_rows:   データ行のリスト（文字列のみ）。空文字セルは空のまま出す。
    col_widths_cm: 列幅[cm]のリスト。省略時は合計15cmを等分する。
    repeat_header: True ならヘッダ行に w:tblHeader を付ける（改ページ時繰り返し）。
    """
    # 列数はヘッダ1行目の colspan 合計から決める（ヘッダ無しならデータ行から）
    if header_rows:
        ncols = _grid_width(header_rows[0])
    else:
        ncols = len(data_rows[0])
    nrows = len(header_rows) + len(data_rows)

    if col_widths_cm is None:
        col_widths_cm = [15.0 / ncols] * ncols

    table = doc.add_table(rows=nrows, cols=ncols)

    # 表全体の設定: 中央配置・固定レイアウト・罫線（過去テーマ5の表4と同じ）
    tblpr = table._tbl.tblPr
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblpr.append(jc)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    tblpr.append(_make_borders())

    # 列幅（gridCol と各セルの tcW の両方に入れる。マージ前に設定しておくと
    # python-docx がマージ時に幅を合算してくれる）
    for i in range(ncols):
        table.columns[i].width = Cm(col_widths_cm[i])
        for r in range(nrows):
            table.cell(r, i).width = Cm(col_widths_cm[i])

    # ヘッダ行: マージしながら埋める。grid はマージ済みマスの管理用
    grid = [[False] * ncols for _ in range(nrows)]
    for r, row_spec in enumerate(header_rows):
        c = 0
        for spec in row_spec:
            while c < ncols and grid[r][c]:
                c += 1
            if isinstance(spec, dict):
                text = spec.get("text", "")
                colspan = int(spec.get("colspan", 1))
                rowspan = int(spec.get("rowspan", 1))
            else:
                text = spec
                colspan = 1
                rowspan = 1
            cell = table.cell(r, c)
            if colspan > 1 or rowspan > 1:
                cell = cell.merge(table.cell(r + rowspan - 1, c + colspan - 1))
            _write_cell(cell, text, font_size, bold=True)
            for rr in range(r, r + rowspan):
                for cc in range(c, c + colspan):
                    grid[rr][cc] = True
            c += colspan
        if repeat_header:
            trpr = table.rows[r]._tr.get_or_add_trPr()
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            trpr.append(hdr)

    # データ行: 空文字はそのまま空セルにする（勝手に埋めない）
    for i, row in enumerate(data_rows):
        r = len(header_rows) + i
        for c, text in enumerate(row):
            _write_cell(table.cell(r, c), text, font_size, bold=False)

    return table


def _grid_width(row_spec):
    """ヘッダ1行分の colspan 合計（＝表の列数）を数える。"""
    n = 0
    for spec in row_spec:
        if isinstance(spec, dict):
            n += int(spec.get("colspan", 1))
        else:
            n += 1
    return n


def _make_borders():
    """外枠 sz=8・内側 sz=4 の黒単線の w:tblBorders を作る。"""
    borders = OxmlElement("w:tblBorders")
    for name, sz in [("top", "8"), ("left", "8"), ("bottom", "8"),
                     ("right", "8"), ("insideH", "4"), ("insideV", "4")]:
        el = OxmlElement("w:" + name)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    return borders


def _write_cell(cell, text, font_size, bold):
    """セル1個ぶんの書式（上下左右中央・ＭＳ明朝）と文字を入れる。"""
    tcpr = cell._tc.get_or_add_tcPr()
    valign = OxmlElement("w:vAlign")
    valign.set(qn("w:val"), "center")
    tcpr.append(valign)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    if text != "":
        add_rich(p, text, font_size=font_size, font_jp=FONT_JP, bold=bold)
