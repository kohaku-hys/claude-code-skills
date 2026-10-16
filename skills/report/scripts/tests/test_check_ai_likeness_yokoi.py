"""check_ai_likeness.py の AIK1〜AIK3（資料「AI臭い文章とは何なのか」の3特徴）の回帰テスト。2026-10-01 新設。

usage:
    py -3 ~/.claude/skills/report/scripts/tests/test_check_ai_likeness_yokoi.py        （全部）
    py -3 ~/.claude/skills/report/scripts/tests/test_check_ai_likeness_yokoi.py -v     （1件ずつ表示）

出典: Kiminori Yokoi「AI臭い文章とは何なのか」（豊洲会 2026-09-28）
      https://speakerdeck.com/nasuvitz/ai-kusai-bunshou-toha-nanina-no-ka
      当たる例・当たらない例の多くは yomiyasu（https://github.com/nanaism/yomiyasu ・MIT・Copyright (c) 2026 nanaism）
      の取り込み調査で確かめた物。実験レポートで正当な書き方（前者・後者、ON側に倒す、表紙の「― 本レポート ―」等）が
      当たらないことも確かめる。
試験用の docx はその場で python-docx で作り、一時フォルダに置く。本物のレポートは読まない。

2026-10-01 の2回目の直しで足した物:
  ・試験集 ~/.claude/skills/humanize/scripts/tests/ai_smell_cases.json（作者の環境の別スキルにある、3つの検査器の共通の試験集。
    このリポジトリには含まれない。無ければ CorpusCases は skip になる）を読み込み、
    report の期待が n/a でない全件（302件）が期待どおりかを確かめる（CorpusCases）。
    判定: 種別の規則（見つけた物の 'rule'）が check_text に出れば hit、check_reference_text に出れば ref、どちらでもなければ miss。
    比喩の種別（B1〜B13・metaphor_other）は、比喩の規則のどれかが出れば hit とする
  ・ワークスペースの docx・本人の文で回帰を取って見つけた誤検知・取りこぼしの試験（Wf2Regressions）

2026-10-01 の3回目の直しで足した物:
  ・試験集の miss・ref の件で、種別の外の AIK 規則も NG を出さないこと（CorpusCases.test_miss_and_ref_cases_have_no_other_ng。
    許す件は ALLOWED_OTHER_NG に理由つき）
  ・反証の指摘24件（low も含む）の文の試験（Wf3Regressions）。教員の資料の文（近似の効かない領域・電子回路とは：・
    次のコマンドを入力する：・参考図書の副題・この処理を、二つのタスクが行うと…・正解にたどり着けない）を含む

2026-10-01 の4回目の直しで足した物:
  ・試験集に足した290件のうち実験レポートで落ちていた35件を通した（試験集の件数は report が n/a でない件 523件）。
    期待を争う件の一覧 CORPUS_DISPUTED（スライドの試験 test_deck_doctor.py と同じ形 {id: (提案する期待, 理由)}）を作った。
    今は争う件は無い（空）
  ・回帰（本人の実験レポート18本・本人以外の学術文88本・ワークスペースの docx・過去問と教員の PDF・学内 PDF）で
    人の文に当たりかけた物と、原理の境目の試験（Wf4Regressions）

2026-10-01 の5回目（裁定の後）:
  ・裁定は「直す件なし・抽象の語の一覧は変えない」だったので、検査器は変えていない。CORPUS_DISPUTED は空のまま
  ・既知の限界の一覧 KNOWN_LIMITS（{id: 理由}。試験集の件数の3%以下）を作った。今は該当する件は無い（空）

2026-10-01 の6回目（本人の判断と、今日の前の版との比べ）:
  ・本人の判断「① 参考の扱いにする」（2026-10-01 21:04）: 落とし込む・架け橋・刺さる は参考層。試験集の metaphor_other-hit-08・
    hit-12・miss-43〜45 の report の期待は ref に直っている。2回目の試験の「学びを現場に落とし込む。」も参考層で確かめる形に直した
  ・今日の前の版（AIK の規則が無い版 md5 cdeccec4）と比べて、人の文（教員の資料・kako・本人の文）で増えた当たりのうち誤検知の文を
    短くして固めた（Wf6Regressions。⚫ の箇条書き・式のマイナスのダッシュ・式が抜けた行末のコロン・参考層の「まず。」等）

2026-10-01 レポートの2回目（6回目の直しの後の、今日の前の版との比べ）:
  ・スキャンの PDF の OCR の崩れの行のダッシュ6件（教員の再試験・事前演習プリント・本人の手書きノート）を外した。
    その文を短くして固め、人が打つ半角の記号（ファイル名の _・引用符・点だけ・取り消し線・Windows のパス）の行は当たったままで
    あることも確かめる（Wf7Regressions）
"""
import os
import sys
import json
import shutil
import tempfile
import unittest
import subprocess

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)

from docx import Document                      # noqa: E402

import check_ai_likeness as cal                 # noqa: E402


def aik(text):
    """本文テキストの当たりのうち AIK の物だけ (pid, match) で返す（合否に数える層）"""
    return [(f['pid'], f['match']) for f in cal.check_text(text) if f['pid'].startswith('AIK')]


def ref(text):
    """参考層（合否に数えない）の当たりを (pid, match) で返す"""
    return [(f['pid'], f['match']) for f in cal.check_reference_text(text)]


def pids(text):
    return [f['pid'] for f in cal.check_text(text)]


# ------------------------------------------------------------
# 資料の具体例（です・ます は である調 に直した）と修正例
# ------------------------------------------------------------

class MaterialExamples(unittest.TestCase):
    def test_katahou_is_aik2(self):
        # 片方・もう片方・壊れる を読み手に解釈させる
        self.assertIn(('AIK2', 'もう片方'), aik('片方だけを見て基準を作ると、もう片方が壊れる。'))

    def test_swapped_repeat_is_aik3(self):
        self.assertIn(('AIK3', '語を入れ替えた反復'), aik('読んで身に付けてから、効く。効いてから、身に付く。'))

    def test_object_comma_is_aik3(self):
        hits = aik('前提を、経路が代わりに添える。')
        self.assertTrue([h for h in hits if h[0] == 'AIK3' and h[1].startswith('前提を、')], hits)

    def test_negative_lead_is_reference_only(self):
        # 短い否定文→短い肯定文は、学術文で正しいこともあるので参考層だけ
        t = '依存構造は分割できない。動かしながら引き返す。'
        self.assertEqual(aik(t), [])
        self.assertEqual(ref(t), [('AIK1', '否定の前振り')])

    def test_side_tao_follows_what_is_tilted(self):
        # 2026-10-01 の2回目の直しで期待を変えた（1回目は「ON側に倒す」とぶつかるので規則を入れず、資料の例も [] にしていた）。
        # 比喩の原理で、倒す物（〜を・〜は）が書いてあって物なら文字どおり、無いか判断する物なら判断を方向で言う比喩として拾う。
        # 試験集 ai_smell_cases.json の B5-hit-02 も report=hit
        self.assertIn(('AIK2', '側に倒す'), aik('判断に迷うものは、残さない側に倒す。'))
        self.assertEqual(aik('トグルスイッチをON側に倒す。'), [])
        self.assertEqual(aik('スイッチは右側に倒す。'), [])

    def test_fixed_examples_are_clean(self):
        for t in ['開発速度だけを基準にすると、保守性の要件を満たせなくなるおそれがある。',
                  '採否を判断できない項目は、原則として除外する。',
                  '依存関係を分離できないため、稼働中のシステムを変更し、問題が発生した場合は元の状態に戻す。',
                  '実際の業務で使い、効果を確かめながら使い方を習得する。',
                  '運用者は問い合わせ先に応じた前提条件を参照する。']:
            self.assertEqual(aik(t), [], t)
            self.assertEqual(ref(t), [], t)


# ------------------------------------------------------------
# AIK1 (1) 主文を直接書かない
# ------------------------------------------------------------

class Aik1(unittest.TestCase):
    def test_lead_in_phrases(self):
        for t in ['ここで重要なのは、配線の長さである。', '結論から言うと、原因は配線だった。',
                  '注目すべき点は、位相のずれである。', '最も大切なのは、測定前の校正である。',
                  '測定した。一言で言うと、遅い。']:
            self.assertTrue([h for h in aik(t) if h[0] == 'AIK1'], t)

    def test_lead_in_not_at_sentence_start(self):
        for t in ['抵抗が重要な理由は2つある。', '重要な部品を交換した。', '一番上の段に置いた。',
                  'この実験で重要な量は時定数である。']:
            self.assertEqual(aik(t), [], t)

    def test_tannaru_denial(self):
        self.assertTrue([h for h in aik('これは単なる測定誤差ではない。') if h[0] == 'AIK1'])
        self.assertTrue([h for h in aik('速度だけの話ではない。') if h[0] == 'AIK1'])
        self.assertEqual(aik('単なる誤差と判断した。'), [])

    def test_colon_ending_a_sentence(self):
        for t in ['メリットは次の3点です：', '使用した機器を以下に示す：', '手順は以下の通り：']:
            self.assertTrue([h for h in aik(t) if h[0] == 'AIK1'], t)

    def test_colon_labels_and_code_are_not_flagged(self):
        # 名詞の見出し・記入欄（先輩の見本レポートにある書き方）と、コードの行
        body = '\n'.join(['日時：10月1日', '時刻は10:30', '注意：', 'サンプリング定理：', '学籍番号：氏名：',
                          'for i in range(10):', '20: ', '91: <p>1つだけ点灯：',
                          'mcp4725.c: In function ‘main’:'])
        self.assertEqual(aik(body), [])


# ------------------------------------------------------------
# AIK2 (2) 抽象語・比喩の動詞
# ------------------------------------------------------------

class Aik2(unittest.TestCase):
    CASES = [
        # (当たる例, 当たらない例)
        (['キャッシュが切れると、データが静かに壊れる。', '重複した処理は黙ってスキップする。'],
         ['試料を実験台に静かに置いた。', '結果が出るまで黙って待った。']),
        (['デバッグに丸一日溶かした。', '設定に3時間溶かした。', 'デバッグで丸一日時間を溶かした。'],
         ['ハンダを溶かし、配線を接続した。', '湯煎でバターを溶かす時間を計った。']),
        (['原因を1つずつ潰していく。', '残ったバグを潰す。'],
         ['空き缶を潰して捨てた。']),
        (['この設定が地味に効いてくる。', 'あとから効いてくるのが配線の抵抗だ。'],
         ['鎮痛薬を飲むと30分ほどで効いた。', '冷房が効きすぎている。']),
        (['ここを読み違えると前提が崩れる。', 'この一文で設計が一気に崩れる。'],
         ['入力を大きくすると出力波形が崩れた。', '地震で塀が崩れた。']),
        (['この仕組みを深掘りする。'],
         ['穴を深く掘った。']),
    ]

    def test_metaphor_verbs(self):
        for hits, misses in self.CASES:
            for t in hits:
                self.assertTrue([h for h in aik(t) if h[0] == 'AIK2'], t)
            for t in misses:
                self.assertEqual(aik(t), [], t)

    def test_katahou_with_a_name_is_not_flagged(self):
        # 何の片方かが書いてあれば拾わない。前者・後者も拾わない
        for t in ['プローブの片方をGNDに、もう片方を入力につないだ。',
                  '片方の端子をGNDに、もう片方をVccにつないだ。',
                  '前者は直列回路、後者は並列回路である。',
                  '一方、ゲルマニウムは逆方向にも電流が流れた。']:
            self.assertEqual(aik(t), [], t)


# ------------------------------------------------------------
# AIK3 (3) 短文・対句・不自然な読点・記号で印象づける
# ------------------------------------------------------------

class Aik3(unittest.TestCase):
    def test_labeling_de_aru(self):
        self.assertTrue([h for h in aik('これは設計の失敗に他ならない。') if h[0] == 'AIK3'])
        self.assertTrue([h for h in aik('だから電流が減るというわけである。') if h[0] == 'AIK3'])
        self.assertEqual(aik('常に成り立つというわけではない。'), [])

    def test_desu_masu_labeling_is_counted_once_by_p012(self):
        # です・ます形は P-012（です。・ません。）が拾う。AIK3 で二重に数えない
        for t in ['これは失敗に他なりません。', '電流が減るというわけです。']:
            self.assertEqual(aik(t), [], t)
            self.assertIn('P-012', pids(t), t)

    def test_emoji(self):
        for t in ['✅ 完了', '🚀 リリースした']:
            self.assertTrue([h for h in aik(t) if h[0] == 'AIK3'], t)
        for t in ['★印は必須項目', '✓ 確認済み', '①最初の手順', '抵抗は10 kΩ、温度は25 ℃とした。']:
            self.assertEqual(aik(t), [], t)

    def test_object_comma_needs_subject_after(self):
        for t in ['私は、実験を行った。', '電圧を測り、記録した。',
                  '1.電流を、オームの法則により、電圧で表す。', 'これらを、それぞれ、図2～図6に示す。']:
            self.assertEqual(aik(t), [], t)

    def test_dash_in_sentence(self):
        self.assertEqual(aik('電圧は上がる—ただし電流は変わらない。'), [('AIK3', '—')])
        # 「——」は1行1件
        self.assertEqual(aik('3つの量——電圧、電流、抵抗——を測った。'), [('AIK3', '—')])

    def test_dash_like_marks_that_are_not_flagged(self):
        body = '\n'.join(['10–100 Hz の範囲で測った。',                   # 数値範囲の – (U+2013)
                          'テーマ12　LED 制御　 ― 本レポート ―',          # 表紙の副題
                          '― プレレポート ―',
                          '加算平均フィルタ―と移動平均フィルタ―を使う。',  # 長音「ー」の打ち間違い
                          'Vc=Vcc― RcIc',                                 # マイナスの代わり
                          '←―→',                                          # 矢印
                          'I－V特性を測った。', 'コンピューターを使った。'])
        self.assertEqual(aik(body), [])

    def test_swapped_repeat_needs_all_words_shared(self):
        for t in ['理論値と測定値はほぼ一致した。測定値と理論値の差は2%であった。',
                  '抵抗を大きくすると電流が小さくなる。電流が小さくなると発熱も減る。',
                  '温度が上がると電流が増え、その電流でさらに温度が上がる。']:
            self.assertEqual([h for h in aik(t) if h[1] == '語を入れ替えた反復'], [], t)

    def test_swapped_repeat_chain_links_are_not_flagged(self):
        # 2026-10-01 反証で見つかった誤検知。前の文の動詞を次の文で「〜した＋名詞」として受け継ぐ「しりとり」型のつなぎは、
        # 対句ではない。1件目は 2022 年の先輩の実験レポート（1年の基礎実験）の文そのまま
        for t in ['iの値をコピーします。コピーした値に１を加算していきます。加算した値をiにコピーします。',
                  'iの値をコピーする。コピーした値に1を加算する。加算した値をiにコピーする。',
                  'コピーした値に1を加算する。加算した値をiにコピーする。',     # 巡回してもしりとりなら拾わない
                  '波形を観測した。観測した波形を図5に示す。',
                  '電流を測定した。測定した電流を表2に示す。',
                  '回路を製作した。製作した回路を図2に示す。',
                  'グラフを作成した。作成したグラフを図6に示す。',
                  '温度を上げた。上げた温度を記録した。',
                  '信号を増幅した。増幅された信号を観測した。',
                  '抵抗に電圧をかけた。電圧を抵抗で割った。']:
            self.assertEqual([h for h in aik(t) if h[1] == '語を入れ替えた反復'], [], t)

    def test_swapped_repeat_still_hits_true_chiasmus(self):
        # しりとりの外し方は「〜した＋名詞」の形だけなので、対句は拾ったまま
        for t in ['測るから分かる。分かるから測る。',
                  '理解して実装する。実装して理解する。',
                  '書いたコードが人を育てる。育てた人がコードを書く。']:
            self.assertIn(('AIK3', '語を入れ替えた反復'), aik(t), t)

    def test_dash_in_code_lines_is_not_flagged(self):
        # 2026-10-01 反証で見つかった誤検知。授業のプログラムリスト（テーマ8）と、本人の個人レポートのコードのコメント。
        # 行末のコロンと同じく、コードの行（「11:」「//」等）はダッシュも見ない
        body = '\n'.join(['11:// 点灯ボタンがクリックされたとき――25ピンに3.3Vを出力',
                          '// 点灯ボタンがクリックされたとき――25 番ピンに3.3V を出力',
                          '19| // 担当：山田 ―― スコア・ライフ・フェーズを最初に戻し、各パートをそろえる'])
        self.assertEqual(aik(body), [])
        # テーマ8 の実験解説が引く PHP マニュアルの見出し（行頭が関数名＋ダッシュ）も、コードの説明の引用なので拾わない
        self.assertEqual(aik('exec ' + chr(0x2014) + ' 外部プログラムを実行する'), [])
        self.assertEqual(aik('exec() ' + chr(0x2014) + ' 外部プログラムを実行する'), [])

    def test_em_dash_after_katakana_or_number_is_flagged(self):
        # 2026-10-01 反証で見つかった見逃し。打ち間違いの外し方（直前がカタカナ・英数字）は「―」(U+2015) 1本だけに掛ける。
        # AI がよく打つ「—」(U+2014) と、2本続きの「――」(U+2015) は数える
        for t in ['原因はコンデンサ' + chr(0x2014) + '容量の誤差だった。',
                  'オシロスコープ' + chr(0x2014) * 2 + 'これが今回の主役である。',
                  '出力は10 V' + chr(0x2014) + '想定の2倍だった。',
                  'オシロスコープ' + chr(0x2015) * 2 + 'これが今回の主役である。']:
            self.assertEqual(len([h for h in aik(t) if h[0] == 'AIK3']), 1, t)
        # 図の矢印は「—」で書いても拾わない
        self.assertEqual(aik('←' + chr(0x2014) * 2 + '→'), [])
        # 前後とも英字なら英語の文の句読点なので拾わない（人が書いた英語の見出し。授業の演習プリント・英語研修のフレーズ集）
        for t in ['English Conversation ' + chr(0x2014) + ' Quick Reference',
                  'Exercise Sheet ' + chr(0x2014) + ' AIを使いこなすための3つの問い']:
            self.assertEqual(aik(t), [], t)


# ------------------------------------------------------------
# 実験レポートで正当な語・書き方（担当の指示にある物）
# ------------------------------------------------------------

class ReportWordsAreClean(unittest.TestCase):
    def test_report_words(self):
        for t in ['前者は直列、後者は並列である。', 'スイッチを閉じた瞬間に電流が流れた。',
                  'トグルスイッチをON側に倒す。', '真理値表を表2に示す。', '結晶の中を電子が動く。',
                  '水蒸気が凝縮して水滴になった。', '装置の電源を入れた。', '発生した熱量を計算した。',
                  '電流は抵抗で制限される。', 'LEDは点灯しなかった。原因は配線であった。']:
            self.assertEqual(aik(t), [], t)
            self.assertEqual(ref(t), [], t)


# ------------------------------------------------------------
# 参考層（否定の前振り）
# ------------------------------------------------------------

class ReferenceLayer(unittest.TestCase):
    def test_not_negations_are_skipped(self):
        for t in ['誤差は0.4 %ほどしかない。山と谷が同じ周波数に出る。',   # しかない（過去レポートの文の形）
                  'この順で測らなければならない。表1に示す。',            # 〜なければならない
                  '「じゃあ在庫から回せません？」と聞いた。',              # 問いかけ
                  'LEDは点灯しなかった。原因は配線であった。']:          # 過去の観察
            self.assertEqual(ref(t), [], t)

    def test_de_wa_nai_de_aru_is_reference_not_ng(self):
        # 学術文の「〜ではない。〜である。」は、表示するが合否には数えない
        t = 'この値は測定誤差ではない。配線の抵抗による差である。'
        self.assertEqual(aik(t), [])
        self.assertEqual(ref(t), [('AIK1', '否定の前振り')])


# ------------------------------------------------------------
# docx からの読み取りと、コマンドの終了コード
# ------------------------------------------------------------

class DocxAndCli(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='report_aik_')

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def make(self, name, paragraphs):
        doc = Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        path = os.path.join(self.tmp, name)
        doc.save(path)
        return path

    def run_cli(self, path):
        return subprocess.run([sys.executable, os.path.join(SCRIPTS, 'check_ai_likeness.py'), path],
                              capture_output=True, text=True, encoding='utf-8', errors='replace')

    def test_check_reads_docx(self):
        p = self.make('a.docx', ['1. 目的', '読んで身に付けてから、効く。効いてから、身に付く。'])
        self.assertIn('AIK3', [f['pid'] for f in cal.check(p)])
        q = self.make('b.docx', ['依存構造は分割できない。動かしながら引き返す。'])
        self.assertEqual(cal.check(q), [])
        self.assertEqual(len(cal.check_reference(q)), 1)

    def test_program_listing_paragraphs_are_clean(self):
        # 本人のテーマ8本レポートのように、プログラムリストを本文の段落で載せた docx（授業の公式リストの写し）
        p = self.make('list.docx', ['14: <?php', '15:     $pins = array(23, 24, 25);',
                                    '11:// 点灯ボタンがクリックされたとき――25ピンに3.3Vを出力'])
        self.assertEqual([f for f in cal.check(p) if f['pid'].startswith('AIK')], [])

    def test_cli_reference_only_exits_0(self):
        r = self.run_cli(self.make('ref.docx', ['この値は測定誤差ではない。配線の抵抗による差である。']))
        self.assertEqual(r.returncode, 0, r.stdout)
        self.assertIn('[OK]', r.stdout)
        self.assertIn('[参考]', r.stdout)

    def test_cli_aik_hit_exits_1(self):
        r = self.run_cli(self.make('ng.docx', ['ここで重要なのは、配線の長さである。']))
        self.assertEqual(r.returncode, 1, r.stdout)
        self.assertIn('AIK1', r.stdout)

    def test_cli_clean_report_exits_0_without_reference(self):
        r = self.run_cli(self.make('ok.docx', ['テーマ5　CR回路の正弦波入力に対する応答', '― 本レポート ―',
                                               '1. 目的', 'CR回路の周波数特性を測定する。',
                                               '前者は直列回路、後者は並列回路である。']))
        self.assertEqual(r.returncode, 0, r.stdout)
        self.assertNotIn('[参考]', r.stdout)


# ------------------------------------------------------------
# 試験集 ai_smell_cases.json（3つの検査器の共通の試験集）の実験レポートの期待
# ------------------------------------------------------------

CASES_PATH = os.path.join(os.path.expanduser('~'), '.claude', 'skills', 'humanize', 'scripts', 'tests', 'ai_smell_cases.json')
METAPHOR_FAMILIES = ('B1', 'B2', 'B3', 'B4', 'B5', 'B6', 'B12', 'B13', 'metaphor_other')


def report_level(family, text):
    """試験集の report の判定（hit / ref / miss）を、見つけた物の 'rule' から決める"""
    rules = METAPHOR_FAMILIES if family in METAPHOR_FAMILIES else (family,)
    if [f for f in cal.check_text(text) if f['rule'] in rules]:
        return 'hit'
    if [f for f in cal.check_reference_text(text) if f['rule'] in rules]:
        return 'ref'
    return 'miss'


# report の期待が miss・ref の件で、種別の外の AIK 規則が NG（exit 1）を出してよい件と、その規則。
#   neg_then_pos-miss-29「設計の失敗に他ならない。次の版で直す。」は、否定の前振りではない（参考層は出ない）が、
#   1文目の「に他ならない」は名札の決め台詞（kimezerifu）として NG になるのが正しい
ALLOWED_OTHER_NG = {
    'neg_then_pos-miss-29': ('kimezerifu',),
}

# 試験集の report の期待が間違っていると考える件（スライドの試験 test_deck_doctor.py の CORPUS_DISPUTED と同じ形）。
#   {id: (提案する期待 'hit'・'ref'・'miss', 理由)}。ここに書いた件は、提案する期待で確かめる。裁定役が試験集を直したら消す。
#   2026-10-01 4回目の時点で争う件は無い（35件とも、人の文の当たりを増やさずに試験集の期待どおりに通った）。
#   5回目（裁定の後）も空のまま。裁定は「直す件なし・語の一覧は変えない」だった
CORPUS_DISPUTED = {
}

# 既知の限界: 直すと人の文の誤検知が増えるので、試験集の期待どおりには直さない件。{id: 理由}。
#   ここに書いた件は試験集の期待を確かめない（理由は報告の not_fixed にも書く）。
#   試験集の件数の3%以下に収める（test_known_limits_are_few）。2026-10-01 5回目の時点で該当する件は無い（空）
KNOWN_LIMITS = {
}


def expected_report(case):
    """試験集の件の report の期待（争う件は CORPUS_DISPUTED の提案）"""
    if case['id'] in CORPUS_DISPUTED:
        return CORPUS_DISPUTED[case['id']][0]
    return case['expect']['report']


class CorpusCases(unittest.TestCase):
    def test_every_report_case(self):
        if not os.path.exists(CASES_PATH):
            self.skipTest('試験集が無い: ' + CASES_PATH)
        with open(CASES_PATH, encoding='utf-8') as f:
            cases = json.load(f)['cases']
        checked = 0
        for c in cases:
            want = expected_report(c)
            if want == 'n/a':
                continue
            checked += 1
            if c['id'] in KNOWN_LIMITS:
                continue
            with self.subTest(id=c['id'], text=c['text']):
                self.assertEqual(report_level(c['family'], c['text']), want, c.get('note', ''))
        # 2026-10-01 の時点で report が n/a でない件は302件、4回目の時点（試験集 789件）で523件。
        # 試験集が減って試験が空振りしていないかを見る
        self.assertGreaterEqual(checked, 520)

    def test_disputed_ids_exist(self):
        # 争う件の id が試験集にあること（試験集の id を付け直した時に、効かない行が残らないように）
        if not os.path.exists(CASES_PATH):
            self.skipTest('試験集が無い: ' + CASES_PATH)
        with open(CASES_PATH, encoding='utf-8') as f:
            ids = set(c['id'] for c in json.load(f)['cases'])
        for cid, (want, why) in CORPUS_DISPUTED.items():
            self.assertIn(cid, ids)
            self.assertIn(want, ('hit', 'ref', 'miss'))
            self.assertTrue(why)

    def test_known_limits_are_few(self):
        # 既知の限界の id が試験集にあり、理由が書いてあり、件数が試験集の3%以下であること
        if not os.path.exists(CASES_PATH):
            self.skipTest('試験集が無い: ' + CASES_PATH)
        with open(CASES_PATH, encoding='utf-8') as f:
            ids = set(c['id'] for c in json.load(f)['cases'])
        for cid, why in KNOWN_LIMITS.items():
            self.assertIn(cid, ids)
            self.assertTrue(why)
        self.assertLessEqual(len(KNOWN_LIMITS), len(ids) * 0.03)

    def test_miss_and_ref_cases_have_no_other_ng(self):
        # 2026-10-01 3回目（反証の指摘）: miss・ref の件で、種別の外の AIK 規則の誤検知（exit 1）が出ても上の試験は通る。
        # どの AIK 規則でも NG にならないことを確かめる（許す件は ALLOWED_OTHER_NG に理由つきで書く）
        if not os.path.exists(CASES_PATH):
            self.skipTest('試験集が無い: ' + CASES_PATH)
        with open(CASES_PATH, encoding='utf-8') as f:
            cases = json.load(f)['cases']
        for c in cases:
            if expected_report(c) not in ('miss', 'ref'):
                continue
            if c['id'] in KNOWN_LIMITS:
                continue
            allowed = ALLOWED_OTHER_NG.get(c['id'], ())
            with self.subTest(id=c['id'], text=c['text']):
                ng = [f['rule'] for f in cal.check_text(c['text']) if f['pid'].startswith('AIK') and f['rule'] not in allowed]
                self.assertEqual(ng, [], c.get('note', ''))

    def test_findings_have_rule_names(self):
        # 見つけた物には規則の名前（'rule'）が付く。P- と SDR は pid と同じ
        for f in cal.check_text('ここで重要なのは、配線の長さである。回路における電流を測る。'):
            self.assertIn('rule', f)
        rules = [f['rule'] for f in cal.check_text('ここで重要なのは、配線の長さである。回路における電流を測る。')]
        self.assertIn('preface', rules)
        self.assertIn('P-022', rules)


# ------------------------------------------------------------
# 2回目の直しの回帰（ワークスペースの docx 577本・本人の文・教員の文で見つけた物）
# ------------------------------------------------------------

def rules_of(text):
    return [f['rule'] for f in cal.check_text(text) if f['pid'].startswith('AIK')]


def ref_rules_of(text):
    return [f['rule'] for f in cal.check_reference_text(text)]


class Wf2Regressions(unittest.TestCase):
    def test_literal_metaphor_verbs_are_not_flagged(self):
        # 人・物・薬・物質・建物が主語や目的語の文字どおりの文（課題文に挙がった物と、回帰で見つけた物）
        for t in ['祖父は静かに死んだ。', '母は私の漫画を黙って捨てた。', '水深10mまで深く潜った。',
                  '鎮痛薬は後から効いてきた。', '試料を塩酸で24時間溶かした。', 'ハンダを長時間溶かすとランドがはがれる。',
                  'はんだを長時間溶かすとランドがはがれる。', '空き缶を1つずつ潰して回収した。', '町家の土台が腐って崩れていた。',
                  '陶芸家は80歳で新境地を開いた。', '肌感触のよい綿の生地を選んだ。', '布の手触りを比べた。',
                  '1GHzでは集中定数の前提が崩れる。', 'スイッチは右側に倒す。',
                  '時間を潰した。',                                              # 暇つぶしの慣用句
                  '書式が潰さないように注意してください。',                      # 教員の見本の文
                  'プローブを手に持ち、もう片方の手の指に対してプローブ先端を接触させる。',  # 教員の実験書の文
                  '各場面に背景がわかる予備知識を添えた。',                        # 人が情報を書き足す
                  '自分たちで考えたアイデアを実際の動きに落とし込んで、何度も試した。',  # 本人の感想の形
                  '4．第9回の作業工程表に落とし込む準備をする．',                  # 教員の課題文
                  'AI に渡したプロンプトの中で、特に効いた1行はどれか。']:          # 効いた＋名詞（その名詞が主語）
            self.assertEqual(rules_of(t), [], t)

    def test_abstract_metaphor_verbs_are_flagged(self):
        for t, rule in [('もう一つ効いたのは、回答前にレビューを頼むことだ。', 'B4'),
                        ('書いたコードのバグを AI と一緒に潰す', 'B3'),
                        ('ライトな温度感を狙う。', 'B13'),
                        ('WaitTimer(16)を深掘りする。', 'B12'),                  # 関数の呼び出しはプログラムの中の物
                        ('処理（例：WaitTimer(16)）を深掘りする。', 'B12')]:       # かっこ書きを飛ばして「処理」
            self.assertIn(rule, rules_of(t), t)
        # 「学びを現場に落とし込む。」は 2026-10-01 本人の判断「① 参考の扱いにする」で参考層に移した（試験集 metaphor_other-hit-08 の
        # report の期待も ref に直った）。合否に数える層には出ず、参考層に出る
        self.assertNotIn('metaphor_other', rules_of('学びを現場に落とし込む。'))
        self.assertIn('metaphor_other', ref_rules_of('学びを現場に落とし込む。'))

    def test_simply_not_with_content_in_same_sentence(self):
        # 同じ文で中身を書く「〜ではなく、Y」「〜ではない、Y」は拾わない（本人・教員の文）
        for t in ['就職は単なる職に就くことではなく、自分の価値観に合った企業と出会うことだ。',
                  '５Ｓは単なる職場をきれいにする活動だけではない、あらゆるビジネスの基本である。']:
            self.assertEqual(rules_of(t), [], t)
        self.assertIn('simply_not', rules_of('それは単なる偶然ではなかった。'))

    def test_colon_before_formula_is_not_flagged(self):
        # 次の行が式（Word の数式を本文に入れた物も）なら、式の前置き。次の行が文なら拾う
        self.assertEqual(rules_of('両辺の y′ を消すと:\n0 = xy″−12y′y″'), [])
        self.assertEqual(rules_of('A·Ã = |A|·E を |A| で割ると：\nA·{(1/|A|)·Ã} = E。よって右辺が A⁻¹。'), [])
        self.assertIn('colon_end', rules_of('両辺の y′ を消すと:\nこの式を整理する。'))
        # 次の行の式の中の「;」（|a,b; a,b|）でコードの行にしない
        self.assertIn('colon_end', rules_of('定理2.2.3を組み合わせると示せます：\n'
                                            '(2行目を2つの行の和に分け、後者は k・|a,b; a,b| = k·0 = 0)。'))

    def test_reference_layer_skips_human_sentences(self):
        # 本人の文（デイリーレポート・個人レポート・実験レポート・感想・キャリア形成のレポート）
        for t in ['中国語も練習しないといけない。半導体をもっと学びたいなと思った。',
                  '今回は補正を入れていない。ただ、これだと画面の外へ出てしまう。',
                  'なお、半加算器は、入力信号A、Bに対して、SとCの2種類の異なる出力が得られる回路である。',
                  'また、企業選択においては、規模や待遇だけではなく、誇り、やりがい、自分の専門性が活かせるかといった視点に共感した。',
                  'まず、障害者とは、「身体障害、知的障害、精神障害（発達障害を含む。）その他の心身の機能の障害がある者をいう。」'
                  '社会的障壁とは、「事物、制度、慣行、観念その他一切のものをいう。」',
                  'コンデンサは直流を通さない。交流は通す。']:
            self.assertEqual(ref_rules_of(t), [], t)
        self.assertEqual(ref_rules_of('速さは要らない。しかし正確さは要る。'), ['neg_then_pos'])
        self.assertEqual(ref_rules_of('まず、電源を、入れて、少し、待ってから、測る。'), ['many_commas'])

    def test_tokasu_is_decided_by_what_melts(self):
        # 溶ける物（目的語、無ければ主語・主題）が物質なら文字どおり。人の主語（私は）は溶かす人なので数えない
        for t in ['試料は1時間溶けなかった。', '試料は塩酸で24時間溶かした。', '氷は1時間溶けずに残った。']:
            self.assertEqual(rules_of(t), [], t)
        for t in ['私は3時間溶かした。', '僕は環境構築で丸一日溶かした。']:
            self.assertIn('B2', rules_of(t), t)

    def test_dash_leaders_and_marks_are_not_flagged(self):
        # 3本以上の罫線・リーダー、式（〓）・式番号の隣、かっこの中の空欄の印、行の頭の ― 1本（箇条書きの印）
        # （ワークスペースの docx の今の当たりを全部読んで見つけた物。教員の経済学の配布資料・学校のメールの見本・人の議事録）
        for t in ['貿易前　―――　　服地1着　＝　酒1樽', '〓　　―――　①', '〓であるから、〓　――　②',
                  '――――　第1次世界大戦(1914年～1918年)　――――', '————————————————————————————',
                  '生産物市場の需給均衡―〓曲線－', '07_1: 余因子行列とその性質 [—]', '―9時から準備']:
            self.assertEqual(rules_of(t), [], t)
        # AI がよく打つ形は当たったまま（行の頭の「— 副題」、用語の後のダッシュ、式番号でない ① の隣の見出しのダッシュ）
        for t in ['— 話しかけていい人が一目でわかる仕組み', 'バンドギャップ — 価電子帯と伝導帯のエネルギー差。',
                  '3. 午後① — ガラス工芸館と近くの神社（見学）', '半導体展示館 — ①2nmの先、微細化の次の一手は?']:
            self.assertIn('dash', rules_of(t), t)

    def test_colon_before_formula_object_mark(self):
        # 次の行が式・図の物の跡（〓）なら式の前置き（教員の経済学の問題）。次の行が文なら拾う
        self.assertEqual(rules_of('次の経済を考える：\n〓\n〓'), [])
        self.assertIn('colon_end', rules_of('次の経済を考える：\nこの経済は閉鎖経済である。'))

    def test_reference_layer_skips_formula_only_second_sentence(self):
        # 2文目が式・数値だけなら、否定の後の主文ではない（予習ノートの「1 Å = 10⁻¹⁰ m」）
        self.assertEqual(ref_rules_of('面の法線からの角ではない。1 Å = 10⁻¹⁰ m'), [])

    def test_sentence_split_after_closing_quote(self):
        # 「。」」の後で新しい文が始まる時は切る。かっこの後に助詞が続く時と、（…。）の後は切らない
        self.assertEqual(len(cal.split_line('障害者とは、「…をいう。」社会的障壁とは、「…をいう。」')), 2)
        self.assertEqual(len(cal.split_line('「じゃあ在庫から回せません？」と聞いた。')), 1)
        self.assertEqual(len(cal.split_line('精神障害（発達障害を含む。）その他の障害をいう。')), 1)


# ------------------------------------------------------------
# 3回目の直し（2026-10-01 反証の指摘24件。low も全部）
# ------------------------------------------------------------

def level(family, text):
    return report_level(family, text)


class Wf3Regressions(unittest.TestCase):
    def assertLevels(self, cases):
        for family, text, want in cases:
            with self.subTest(text=text):
                self.assertEqual(level(family, text), want)

    def test_spaces_and_greek_do_not_change_the_result(self):
        # 名詞と助詞の間の空白（本人のレポートの「測定値 Rx は」）とギリシャ文字（増幅率β・時定数τ・sinθ）。
        # 空白のある文と無い文が同じ判定になる
        for a, b, family in [('試料 A は塩酸で24時間溶かした。', '試料Aは塩酸で24時間溶かした。', 'B2'),
                             ('トグルスイッチ SW1 は ON 側に倒した。', 'トグルスイッチSW1はON側に倒した。', 'B5'),
                             ('コンデンサ C1 は静かに壊れていた。', 'コンデンサC1は静かに壊れていた。', 'B1'),
                             ('高域では C2 が効く。', '高域ではC2が効く。', 'B4'),
                             ('増幅率 β が効いてくる。', '増幅率βが効いてくる。', 'B4')]:
            with self.subTest(text=a):
                self.assertEqual(level(family, a), level(family, b))
                self.assertEqual(level(family, a), 'miss')
        self.assertLevels([('B5', 'プローブ CH2 は GND 側に倒した。', 'miss'), ('B3', '抵抗 R は過電流で潰した。', 'miss'),
                           ('B4', '時定数 τ が効いてくる。', 'miss'), ('B4', 'だから sinθ が効く。', 'miss')])  # 教員の解説の文

    def test_kiku_relative_clause_and_cleft(self):
        # 「Xの効かない＋名詞」の X（関係節の主格の「の」）と、分裂文「効くのは X」の X が主語。物なら拾わない
        self.assertLevels([
            ('B4', '近似の効かない領域（E − EF ≪ kBT）は、半導体の場合ギャップ内になる。', 'miss'),   # 教員のスライド
            ('B4', '冷房の効かない部屋で測定した。', 'miss'), ('B4', 'ブレーキの効かない台車を使った。', 'miss'),
            ('B4', 'ブレーキの効いた台車を使った。', 'miss'),
            ('B4', '低域遮断周波数に効くのはCEである。', 'miss'),
            ('B4', '低域遮断周波数に最も効くのはCEである。', 'miss'),          # 教員の授業資料の文（「最も」の「も」は助詞でない）
            ('B4', '低域で効いてくるのは結合コンデンサである。', 'miss'),
            ('B4', 'この差に効いているのは配線の抵抗である。', 'miss'),
            # 分裂文の主語が抽象（習慣・〜こと）なら拾ったまま（yomiyasu raw_ai の文）
            ('B4', 'もう一つ効いたのは、レビューで「これ、新人が読んで分かる？」と聞く習慣だ。', 'hit'),
            ('B4', 'もう一つ効いたのは、回答前にレビューを頼むことだ。', 'hit'),
            ('B4', '何度押しても効かなかった。', 'hit')])

    def test_colon_before_command_and_definition_heading(self):
        self.assertLevels([
            ('colon_end', '③次のコマンドを入力する：\ncp -r /home/inst/onomi/ev3rt-beta7-2-release ~/', 'miss'),  # 教員の手順
            ('colon_end', '次のコマンドを実行する：\nsudo raspi-config', 'miss'),
            ('colon_end', '次のコマンドを入力する：\nls -l /home/pi', 'miss'),
            ('colon_end', '次のコマンドで起動する：\npython3 led.py', 'miss'),
            ('colon_end', '次のコマンドを入力する：\n$ ls -l', 'miss'),
            ('colon_end', '使用した機器を以下に示す：\nオシロスコープ', 'hit'),
            # 定義の見出し（教員の資料）
            ('colon_end', '• 電子回路とは：', 'miss'), ('colon_end', '• 二分探索とは：', 'miss'),
            ('colon_end', 'テキストファイルとは：\n・文字コードだけで構成されたファイル．', 'miss'),
            ('colon_end', 'トランジスタとは：\n電流を増幅する素子である。', 'miss'),
            # 名詞に数が付いた見出し。「は・の＋数」の述語は拾ったまま
            ('colon_end', '測定の注意点2つ：\n・プローブの接地を短くする\n・レンジを合わせる', 'miss'),
            ('colon_end', '(2) ハッシュ関数の特徴3つ:\n・速い', 'miss'),
            ('colon_end', 'ポイントは3つ：\n速い・安い・うまい', 'hit'),
            ('colon_end', '使った機器は次の3つ：\nオシロスコープ・電源・テスター', 'hit')])

    def test_connective_colon(self):
        # 行の全体が接続の語＋コロンで、次の行が文なら拾う。次の行が図の名前なら「例：」と同じ見出し（教員の物理の講義資料）
        self.assertLevels([('colon_end', '例えば：\n遮断周波数が5%低かった。', 'hit'),
                           ('colon_end', '要するに：\n遮断周波数が5%低かった。', 'hit'),
                           ('colon_end', 'つまり：\n遮断周波数が5%低かった。', 'hit'),
                           ('colon_end', 'その結果：\n遮断周波数が5%低かった。', 'hit'),
                           ('colon_end', '例えば：\nPからQへの移動\n各地点での速度', 'miss'),
                           ('colon_end', '結論：\n遮断周波数が5%低かった。', 'miss')])

    def test_book_subtitle_dash(self):
        self.assertLevels([
            ('dash', '「アカデミック・スキルズ(第3版) ―大学生のための知的技法入門（慶應義塾大学出版会）」', 'miss'),  # 教員の参考図書
            ('dash', '[1] 藤井信生，『アナログ電子回路―集積回路化時代の―』，昭晃堂，1984．', 'miss'),
            ('dash', '[2] 佐藤望編，『アカデミック・スキルズ（第3版）―大学生のための知的技法入門』，慶應義塾大学出版会，2020．', 'miss'),
            ('dash', '藤井信生：アナログ電子回路―集積回路化時代の―（昭晃堂，1984）', 'miss'),
            ('dash', 'バンドギャップ ' + chr(0x2014) + ' 価電子帯と伝導帯のエネルギー差。', 'hit'),
            ('dash', '遮断周波数' + chr(0x2015) * 2 + 'つまり出力が3 dB下がる周波数' + chr(0x2015) * 2 + 'は1.6 kHzであった。', 'hit')])

    def test_tame_with_time_clause(self):
        self.assertLevels([('tame', 'この処理を、二つのタスクが行うと…', 'miss'),      # 教員のスライド
                           ('tame', 'この処理を、二つのタスクが行うと値が壊れる。', 'miss'),
                           ('tame', '試料を、温度が下がってから取り出した。', 'miss'),
                           ('tame', '電源を、教員が確認してから入れた。', 'miss'),
                           ('tame', '回路を、班員が組み終えてから確かめた。', 'miss'),
                           ('tame', '前提を、経路が代わりに添える。', 'hit'),
                           ('tame', '設計を、運用が支える。', 'hit')])

    def test_things_are_not_abstract(self):
        # 人が主語の B1、答案（紙）、構造（物の構造）、名指しした性質（線形性・対称性）
        self.assertLevels([('B1', '母は私の答案を黙って捨てた。', 'miss'), ('B1', '橋の鉄骨構造が静かに壊れていった。', 'miss'),
                           ('B1', '先輩に挨拶したが黙って無視された。', 'miss'),
                           ('B6', '高温で結晶構造が崩れる。', 'miss'),
                           ('B6', '入力が大きくなると線形性が崩れ、出力がひずむ。', 'miss'),
                           ('B6', '対称性が崩れると、出力にオフセットが出る。', 'miss'),
                           ('B12', '問題の構造を掘り下げる。', 'hit'), ('B3', '改善案を1つずつ潰した。', 'hit'),
                           ('B1', 'データが静かに壊れる。', 'hit'), ('B6', '設計判断が崩れる。', 'hit')])

    def test_swap_with_different_subjects(self):
        # 主語の違う2つの物の働きを並べた技術の文は拾わない。主語そのものを入れ替える対句は拾ったまま
        self.assertLevels([('swap', 'スピーカーは電気信号を音に変える。マイクは音を電気信号に変える。', 'miss'),
                           ('swap', 'モーターは電気を回転に変える。発電機は回転を電気に変える。', 'miss'),
                           ('swap', '送信側は信号を符号に変える。受信側は符号を信号に戻す。', 'miss'),
                           ('swap', '変調器は音声を電波に変える。復調器は電波を音声に戻す。', 'miss'),
                           ('swap', '道具が人を変える。人が道具を選ぶ。', 'hit')])

    def test_swap_in_one_sentence(self):
        # 1文の中の対句（連用中止で切った前後の節）。名詞＋助詞で切れる読点・条件の文は拾わない
        self.assertLevels([('swap', '理論が測定を導き、測定が理論を確かめる。', 'hit'),
                           ('swap', '測定が理論を支え、理論が測定を導く。', 'hit'),
                           ('swap', '抵抗に電圧をかけ、電圧を抵抗で割った。', 'miss'),
                           ('swap', '電圧を上げると電流が増え、電流が増えると発熱も増えた。', 'miss')])

    def test_mou_katahou_with_whole_named(self):
        self.assertLevels([('vague_ref', 'ワニ口クリップを片方だけ外し、もう片方はつないだままにした。', 'miss'),
                           ('vague_ref', '入力端子を片方接地し、もう片方に信号を加えた。', 'miss'),
                           ('vague_ref', '一対の電極を用意し、片方を正極、もう片方を負極とした。', 'miss'),
                           ('vague_ref', 'イヤホンを片方だけ付け、もう片方は外した。', 'miss'),
                           ('vague_ref', '片方だけを見て基準を作ると、もう片方が壊れる。', 'hit')])
        # 「いずれ・はずれ」は「ずれ」ではない
        self.assertEqual(cal.arguments('いずれも', 'も'), [])
        self.assertEqual(cal.noun_before('測定値のずれ'), 'ずれ')

    def test_tokasu_nearest_argument_and_sun(self):
        self.assertLevels([('B2', '朝日を浴びて雪が溶けた。', 'miss'), ('B2', '西日を受けてチョコレートが溶けた。', 'miss'),
                           ('B2', '数日を経て沈殿が溶けた。', 'miss'), ('B2', '冷凍庫から出しても3日間溶けなかった。', 'miss'),
                           ('B2', '気づけば一日が溶けていた。', 'hit'), ('B2', '僕は環境構築で丸一日溶かした。', 'hit')])

    def test_soeru_topic_is_not_the_one_who_adds(self):
        self.assertLevels([('metaphor_other', '測定値は必ず単位を添えて書く。', 'miss'),
                           ('metaphor_other', '数値は単位を添えて書く。', 'miss'),
                           ('metaphor_other', '各測定値は単位を添えて記録した。', 'miss'),
                           ('metaphor_other', '測定結果は図を添えて表2に示す。', 'miss'),
                           ('metaphor_other', '実験データは考察を添えて提出した。', 'miss'),
                           ('metaphor_other', '前提を、経路が代わりに添える。', 'hit')])

    def test_b13_literal_owner_and_landing(self):
        self.assertLevels([('B13', '川魚は泥臭く感じる。', 'miss'), ('B13', '鯉の泥臭さを抜くため水に浸けた。', 'miss'),
                           ('B13', '着地して地に足がついた瞬間に力が最大になった。', 'miss'),
                           ('B13', '泥臭く手を動かした。', 'hit'), ('B13', '地に足のついた提案だ。', 'hit')])

    def test_tsubusu_without_object(self):
        self.assertLevels([('B3', '乳鉢で細かく潰した。', 'miss'), ('B3', '配線ミスで1個潰した。', 'miss'),
                           ('B3', '原因の潰し方を書いた。', 'hit')])

    def test_tadoritsuku_answer(self):
        self.assertLevels([
            ('metaphor_other', 'この解答例では，起電力に任意の初期位相を与えた場合や起電力の初期位相が不明な場合は'
                               '不自然な記述となるため正解にたどり着けない場合がある．', 'miss'),   # 教員の過去問の解説
            ('metaphor_other', '試行錯誤の末、ようやく答えにたどり着いた。', 'miss'),
            ('metaphor_other', 'この実習でよかったのは、自分だけでは出せなかった案に6人でたどり着けたことである。', 'hit')])

    def test_tao_without_object_uses_the_side_word(self):
        self.assertLevels([('B5', 'ON側に倒した状態で、出力電圧を測った。', 'miss'), ('B5', '次に、OFF側に倒した。', 'miss'),
                           ('B5', '右側に倒すとLEDが点灯する。', 'miss'),
                           ('B5', 'このとき、ON側に倒したスイッチが戻った。', 'miss'),
                           ('B5', '手前側に倒したレバーを元に戻した。', 'miss'),
                           ('B5', '迷ったら安全側に倒す。', 'hit'), ('B5', '判断に迷ったら除外側に倒す。', 'hit'),
                           ('B5', '迷ったら厳しめの方に倒す。', 'hit')])

    def test_kuzureru_named_by_quantity(self):
        self.assertLevels([('B6', '抵抗値のバランスが崩れると、ブリッジの出力は0にならない。', 'miss'),
                           ('B6', '周波数が10 kHzを超えると、測定値の傾向が崩れた。', 'miss'),
                           ('B6', '実測値の対称性が崩れた。', 'miss'), ('B6', '出力値の直線性が崩れた。', 'miss'),
                           ('B6', '測定結果の再現性が崩れた。', 'miss'), ('B6', '周波数特性が崩れた。', 'miss'),
                           ('B6', '戦略の根拠が崩れた。', 'hit'), ('B6', '議論の土台が崩れた。', 'hit'),
                           ('B6', '前提が崩れる。', 'hit')])

    def test_relative_clause_head(self):
        self.assertLevels([('B3', '潰した空き缶を回収した。', 'miss'), ('B3', '潰したケーブルは交換した。', 'miss'),
                           ('B1', '静かに壊れたコンデンサを交換した。', 'miss'), ('B12', '深掘りした溝にケーブルを通した。', 'miss'),
                           ('B1', '黙って捨てたデータが後で必要になった。', 'hit'),
                           ('B1', '黙って壊れていた自動の仕組みを4か所直した。', 'hit')])    # Claude のデイリーノートの文

    def test_measurement_concepts_are_abstract(self):
        self.assertLevels([('B3', '誤差を1つずつ潰した。', 'hit'), ('B3', '測定値のばらつきを1つずつ潰した。', 'hit'),
                           ('B12', 'この現象を深掘りした。', 'hit'), ('B12', '位相差の挙動を深掘りした。', 'hit'),
                           ('B12', 'このズレを深掘りした。', 'hit'), ('metaphor_other', '回路の挙動を紐解く。', 'hit'),
                           ('B4', '配線抵抗の影響が後から効いてくる。', 'hit'), ('B4', '配線抵抗の効果が後から効いてくる。', 'hit'),
                           ('B13', '測定の手触りがつかめた。', 'hit'), ('B4', '誤差が静かに効いてくる。', 'hit'),
                           ('B4', '負荷抵抗の差が効いて電流が変わった。', 'miss')])

    def test_preface_words(self):
        self.assertLevels([('preface', '何より大切なのは、温度管理である。', 'hit'),
                           ('preface', '何よりも重要なのは、温度管理である。', 'hit'),
                           ('preface', '興味深いのは、周波数を上げると位相差が小さくなった点である。', 'hit'),
                           ('preface', '正直に言うと、この結果は予想していなかった。', 'hit'),
                           ('preface', '注目したいのは、立ち上がり時間である。', 'hit'),
                           ('preface', '特筆すべきは、立ち上がり時間である。', 'hit'),
                           ('preface', '重要な点は、配線の抵抗である。', 'hit'),
                           ('preface', '興味深いのは分かるが、時間が足りない。', 'miss'),
                           ('preface', '正直に言うことが大切だ。', 'miss')])

    def test_many_commas_after_connective(self):
        self.assertLevels([('many_commas', 'そして、その結果、電流は、確かに、減少した。', 'ref'),
                           ('many_commas', 'その結果、電流は10 mAに減少した。', 'miss')])

    def test_closing_with_to(self):
        self.assertLevels([('closing', '本レポートが今後の学習の一助となれば幸いである。', 'hit'),
                           ('closing', '本レポートが今後の実験の参考となれば幸いである。', 'hit'),
                           ('closing', '本報告が後続の実験の助けとなれば幸いである。', 'hit'),
                           ('closing', '本報告が今後の実験に役立てば幸いである。', 'hit'),
                           ('closing', '本報告が今後の実験の役に立てば幸いである。', 'hit'),
                           ('closing', '参考となる値を表1に示す。', 'miss')])

    def test_noun_metaphors_are_reference_only(self):
        # 名詞の比喩（落とし穴・核心・主役・〜への入口）は参考層だけ。合否（exit）には数えない
        self.assertLevels([('metaphor_other', '位相の遅れという落とし穴があった。', 'ref'),
                           ('metaphor_other', 'この実験の核心は、プローブの補正にある。', 'ref'),
                           ('metaphor_other', '送信側にも落とし穴がある。', 'ref'),
                           ('metaphor_other', '理解への入口になる。', 'ref'),
                           ('metaphor_other', '会場への入口は北側にある。', 'miss'),
                           ('metaphor_other', '公園に落とし穴を掘った。', 'miss')])


# ------------------------------------------------------------
# 4回目の直し（2026-10-01 試験集に足した290件のうち実験レポートで落ちていた35件と、その回帰）
# ------------------------------------------------------------

class Wf4Regressions(unittest.TestCase):
    def assertLevels(self, cases):
        for family, text, want in cases:
            with self.subTest(text=text):
                self.assertEqual(level(family, text), want)

    def test_human_sentences_found_in_regression(self):
        # 直している途中で人の文に当たりかけた物（回帰の母集団で見つけた文そのまま）。どれも当てない
        self.assertLevels([
            ('preface', 'ポイントは、電気抵抗率！', 'miss'),                                  # 教員のスライド（第2回）
            ('B13', 'ア：インクリメンタルモデル（段階的モデル）。', 'miss'),                   # 基本情報の1問1答（開発手法の名前）
            ('metaphor_other', '（コントラスト比，解像度が低い）', 'miss'),                    # 教員のプレゼンの講義資料
            ('metaphor_other', '解像度が高い最近の側面境界値を使うこと', 'miss'),              # 理化学研究所の研究発表の PDF
            ('metaphor_other', '処理時間が増加するため、解像度の高いデータなどの処理に時間がかかる。', 'miss'),  # 先輩の発表資料
            ('metaphor_other', '知見を積み重ねる必要があります。', 'miss'),                     # 理化学研究所の研究発表の PDF
            ('metaphor_other', '丁寧に一個ずつ前提を積み重ねていく。', 'miss'),                 # yomiyasu の human コーパス
            ('metaphor_other', '天体と地球の距離の変化を織り込みつつ、軌道を計算した。', 'miss'),  # 科学史の講義資料の形
            ('neg_then_pos', 'よって、変化しないのである。矩形波は図3に示す。', 'miss'),       # 先輩の実験レポートの形（1文目の のである）
            ('dash', '教師あり学習—回帰問題', 'miss'),                                         # 教員のスライドの見出し
            ('dash', 'イースターパーティ（conversation）―後日詳細を提案', 'miss')])            # 人の議事録

    def test_dash_by_usage(self):
        # ダッシュは字でなく使い方で分ける。表のセルの空欄・2つの量の組の名前・文の無い見出しの区切りは外し、
        # 文の途中の間（PDF の行の切れ目で文の終わりが次の行に回った物も）は拾ったまま
        self.assertLevels([
            ('dash', '| 抵抗 | ' + chr(0x2014) + ' | 10 kΩ |', 'miss'),
            ('dash', '応力' + chr(0x2014) + 'ひずみ曲線を描いた。', 'miss'),
            ('dash', '原因はコンデンサ' + chr(0x2014) + '周波数の誤差だった。', 'hit'),
            ('dash', '速さ' + chr(0x2014) + 'それが二人の関係だった。', 'hit'),
            ('dash', '原因は一つ' + chr(0x2015) + '配線の取り違え', 'hit'),               # 主題がある左側は見出しでなく文
            ('dash', 'キーワード対応で' + chr(0x2014) * 2 + 'KM＝知識共有，SCM＝調達', 'hit'),   # Claude の1問1答の PDF（助詞で終わる左側）
            ('dash', 'あるが、実質「避けるだけ」。 だから' + chr(0x2014) * 2 + '上手い/下手の差が出', 'hit')])  # 前に文の終わり

    def test_kiku_target_and_delay(self):
        # 何に効くかが書いてあれば拾わない。時間がたって効く言い方で何に効くかが無ければ、薬・器具の主語の時だけ外す
        self.assertLevels([
            ('B4', '数なら1か所直せば全体に効く。', 'miss'),
            ('B4', '肩こりには、この習慣が後から効いてくる。', 'miss'),
            ('B4', 'ヒーターは5分後に効いてきた。', 'miss'), ('B4', '鎮痛薬は半年後に効いてきた。', 'miss'),
            ('B4', '冷房が最後に効いてきた。', 'miss'),
            ('B4', '配線の抵抗が後から効いてくる。', 'hit'),
            ('B4', '早めに相談したことが、最後に効いた。', 'hit')])

    def test_quiet_other_verbs(self):
        # 静かに・黙って＋一覧に無い動詞は、すぐ前の主語が情報・処理の物か品質の時だけ
        self.assertLevels([
            ('B1', 'ログが、静かに消える。', 'hit'),
            ('B1', 'モーターが静かに回る。', 'miss'), ('B1', '議論が静かに進んだ。', 'miss'),
            ('B1', '時間が静かに過ぎる。', 'miss'), ('B1', '彼は黙って立ち去った。', 'miss'),
            ('B1', '結果が出るまで黙って待った。', 'miss')])

    def test_noun_metaphor_verbs_need_abstract_subject(self):
        self.assertLevels([
            ('metaphor_other', 'トマトが実を結んだ。', 'miss'), ('metaphor_other', '種が芽生えた。', 'miss'),
            ('metaphor_other', '責任感が芽生えた。', 'miss'),                 # 人の気持ちは抽象の一覧に入れない
            ('metaphor_other', 'アイデアが芽生えた。', 'hit'),
            ('metaphor_other', '先輩が語った。', 'miss'), ('metaphor_other', 'この結果が物語っている。', 'miss'),
            ('metaphor_other', '部屋に踏み込んだ。', 'miss'), ('metaphor_other', 'ペダルを踏み込んだ。', 'miss'),
            ('metaphor_other', 'ブロックを土台として積み上げた。', 'miss'),
            ('metaphor_other', '経験を財産として積み重ねる。', 'hit')])

    def test_resolution_owner(self):
        self.assertLevels([
            ('metaphor_other', 'カメラの設定で解像度を上げて撮影した。', 'miss'),
            ('metaphor_other', '画像の解像度を上げた。', 'miss'),
            ('metaphor_other', '顧客の解像度を上げる。', 'hit')])

    def test_preface_metaphor_nouns_and_manner(self):
        self.assertLevels([
            ('preface', 'カギになるのは、温度管理である。', 'hit'), ('preface', '鍵は、机の上にある。', 'miss'),
            ('preface', '正直なところを聞かせてほしい。', 'miss')])

    def test_neg_then_pos_relation_written(self):
        # 2文目が帰結の接続の語で始まる・評価の否定（異常はない）は参考に出さない。対比の「むしろ・しかし」は出す
        self.assertLevels([
            ('neg_then_pos', '入出力間の電圧は変わらない。つまりハイパスフィルタになる。', 'miss'),
            ('neg_then_pos', 'n=10 → 少ない。だから順位検定にした。', 'miss'),
            ('neg_then_pos', '測定値に異常はない。次の測定に移る。', 'miss'),
            ('neg_then_pos', '精度は悪くならない。むしろ上がる。', 'ref')])

    def test_emphasis_tail_and_question_swap(self):
        self.assertLevels([('B6', '設計自体が崩れた。', 'hit'), ('B6', '塀そのものが崩れた。', 'miss'),
                           ('swap', '何を測るか？測って何をするか？', 'miss')])
        self.assertEqual(cal.noun_before('前提そのもの'), '前提')


# ------------------------------------------------------------
# 6回目の直し（2026-10-01 本人の判断「① 参考の扱いにする」と、今日の前の版 check_ai_likeness_before.py（md5 cdeccec4）との比べ）
#   今日の前の版と比べて、人の文（本人の文・本人の実験レポート・kako・教員の資料・yomiyasu の human）で増えた当たりのうち、
#   誤検知だった物の文を短くして固めた。AI がよく打つ形が当たったままであることも一緒に確かめる
# ------------------------------------------------------------

class Wf6Regressions(unittest.TestCase):
    def assertLevels(self, cases):
        for family, text, want in cases:
            with self.subTest(text=text):
                self.assertEqual(level(family, text), want)

    def test_owner_words_are_reference_only(self):
        # 落とし込む・架け橋・刺さる は本人が自分の文で使う語なので、比喩の型でも参考層（表示だけ）。合否に数える層には出さない
        for t in ['学びを現場に落とし込む。', '提案3案を具体的なアクションレベルまで落とし込む',
                  'アイデアを実際の動きに落とし込んで',                         # 本人の1年の報告
                  '交流の架け橋',                                               # 本人の授業レポート
                  '他国との交流の架け橋になり、グローバル化が進められる。',
                  '講習の説明も刺さった',                                       # 本人の学内研修の感想
                  'うつ状態の学生についての説明も刺さった。',
                  '日常業務のレベルまで落とし込み,具体的な目標を立てる',         # 基本情報の過去問の文
                  'この言葉は学生に刺さった。', '刺さる一言を選んだ。']:
            with self.subTest(text=t):
                self.assertEqual(rules_of(t), [])
                self.assertIn('metaphor_other', ref_rules_of(t))
        # 文字どおりの使い方と、何に書き込むかが決まる使い方と、「」で引いた言葉は表示もしない
        for t in ['4．第9回の作業工程表に落とし込む準備をする．',               # 教員の課題文
                  'ボールを穴に落とし込んだ。', '川に架け橋を建設した。',
                  'プローブがブレッドボードに刺さっていない。', 'とげが指に刺さった。',
                  'マックスシェラーは「愛は貧しい知識から豊かな知識の架け橋」であると言いました。']:
            with self.subTest(text=t):
                self.assertEqual(rules_of(t), [])
                self.assertNotIn('metaphor_other', ref_rules_of(t))

    def test_emoji_used_as_marks(self):
        # 教員の PDF の箇条書きの「⚫」（行の頭）、空白の位置を指す「🔺🔺」だけの行、文字化けの行は絵文字に数えない
        self.assertLevels([
            ('emoji', '⚫ プログラムとプロセスの違い', 'miss'),          # 教員の「中間テスト 解説.pdf」
            ('emoji', '⚫ ', 'miss'),
            ('emoji', '  ⚫ コンピュータの5大装置', 'miss'),
            ('emoji', '半角スペースを入れ忘れないように注意してください\n🔺🔺\n🔺', 'miss'),   # 教員の手順書
            ('emoji', '┉芗炘㍚嫎☔մ仼儖', 'miss'),                       # 市民公開講座の PDF の文字化け（アルメニアの字）
            ('emoji', '⚡㕙', 'miss'),                                     # 同じ PDF（CJK 統合漢字拡張A）
            # AI がよく打つ形は当たったまま（行の頭の絵、文の中の ⚫、色で飾る丸、文の後ろの絵）
            ('emoji', '🚀 起動が速い', 'hit'), ('emoji', '⚡ 起動が速い', 'hit'),
            ('emoji', '状態は⚫だった。', 'hit'), ('emoji', '🔴 未解決：配線の確認', 'hit'),
            ('emoji', '🔺 注意して配線する', 'hit'),
            ('emoji', 'Ziが大きい👍、Zoが小さい👍', 'hit')])            # 教員のスライドでも絵の字は型どおり

    def test_dash_used_as_minus_line_or_bullet(self):
        # 教員の解答・配布物で、式のマイナス・負の数・図の線・区切りの線・箇条書きの印として使ったダッシュは拾わない
        D, H, Z = chr(0x2014), chr(0x2015), chr(0x3000)
        self.assertLevels([
            ('dash', 'U(q)=U(0)' + D + 'W(q)と定義できるので', 'miss'),     # 授業の解答
            ('dash', 'したがって、F= Fxi=' + D + '400i', 'miss'),
            ('dash', 'Ul(n) ' + D + 'U2(n) ' + D + 'U3(n) の和をとる', 'miss'),
            ('dash', 'Vinが' + D + ' 1[V] のとき', 'miss'),                 # 電子回路応用の解答
            ('dash', '図中（' + H + '〇）は端子が接続される電極を表している．', 'miss'),   # 授業の宿題
            ('dash', '半導体\n' + H * 2 + '\n電気伝導度が導体と絶縁体との間の物質', 'miss'),   # 区切りの線だけの行
            ('dash', '・反転演算増幅回路、その増幅度 = ' + H, 'miss'),
            ('dash', D + '学生の自由な発想や関心をもとに、新たな挑戦に取り組むもの', 'miss'),   # 市の募集要項の箇条書き
            ('dash', 'まとめ' + H + ' 今回の達成目標', 'miss'),            # 教員のスライドの題
            ('dash', Z * 2 + 'ＧＰＡ' + Z * 3 + H + Z + '大学生活全体の学び方の結果、学び方の質', 'miss'),  # 教員のテキストの対応の表
            ('dash', D + ' so poor contact was the cause.', 'miss'),       # PDF の行の切れ目で英文の途中から始まった行
            # AI がよく打つ形は当たったまま
            ('dash', '原因はコンデンサ' + D + '容量の誤差だった。', 'hit'),
            ('dash', '出力は10 V' + D + '想定の2倍だった。', 'hit'),
            ('dash', '結果は明快だった' + D + '配線ミスである。', 'hit'),
            ('dash', D + ' 話しかけていい人が一目でわかる仕組み', 'hit'),
            ('dash', 'バンドギャップ ' + D + ' 価電子帯と伝導帯のエネルギー差。', 'hit'),
            ('dash', '・vegista_pdf_v4.pdf ' + D + ' 1頁目のヘッダを削除した提出版です。', 'hit'),   # 名前と説明の英語の書き方
            ('dash', 'TOEIC対策 ' + D + ' ⚠ 渡航中は演習できなかった。', 'hit'),               # 飾りの記号の後ろの文
            ('dash', '・情報処理Ⅱ ' + D + ' アクセス権は先にこちらで扱っている', 'hit')])   # ローマ数字も語の字

    def test_colon_before_missing_formula_or_figure(self):
        # 式・図が文字にならずに抜けた配布物（次の行に日本語の字が無い・文の終わりの無い等号の行・「次式で」と名指し）は拾わない
        Z = chr(0x3000)
        self.assertLevels([
            ('colon_end', 'を矢印で示す：\nＺ', 'miss'),                          # 教員の演習ドリル
            ('colon_end', '・抵抗の電圧Vと電流Iは比例する：\nｖＩｖｌ', 'miss'),
            ('colon_end', '次の経済を考える：\n' + Z * 2 + 'Y = C + I', 'miss'),   # 授業の配布物（式だけの行）
            ('colon_end', '次の経済を考える：\n(1)', 'miss'),
            ('colon_end', 'A,B間の電圧を求めると：\n必=(A,Bの開放電圧)=60x40', 'miss'),
            ('colon_end', 'CPUからメモリへの有効アクセス時間Tは次式で定義される：\nh: キャッシュのヒット率', 'miss'),
            # 次の行が文・日本語の語句なら拾ったまま。次の行が無い（文書の終わり）時も拾ったまま
            ('colon_end', '次の経済を考える：\nこの経済は閉鎖経済である。', 'hit'),
            ('colon_end', '使用した機器を以下に示す：\nオシロスコープ', 'hit'),
            ('colon_end', '次の3つを使った：\nテスター・電源・オシロ', 'hit'),
            ('colon_end', '使用した機器を以下に示す：', 'hit')])

    def test_reference_layer_on_teachers_text(self):
        # 参考層の、教員の資料での誤検知（否定でない「まず・はず・過ぎません」、かっこで閉じた否定の2文目、
        # PDF の行の切れ目の「ならない。」、打ち消しの頼み「〜しないでください」、同じ語尾で並べた読点）。資料の例の形は参考に出たまま
        for t in ['調整する必要がある。まず。プローブを接続する。',               # 教員の実験書（「まず、」の打ち間違い）
                  '大きさは2～3%程度に過ぎません。図5に示す。',                   # 教員の講義資料（すぎない の です・ます形）
                  '上がるはず。直せる問題だという証拠。',
                  '（時間差はない。原因と結果という因果関係もない。）',             # 授業の配布物（2文とも否定）
                  'らない。（必ずしもC=x0ではない）',
                  'ならない。休止期間を変更するときも、同様とする。',               # 法令の PDF の行の切れ目
                  '責任を負いません。パスワードを記載しないでください。',           # 教員の小テストの注意書き（2文とも打ち消し）
                  '継続的に、計画的に、安定的に、自己調整をしながら努力できる人かどうかを知る手がかりにもなります。']:
            with self.subTest(text=t):
                self.assertEqual(ref_rules_of(t), [])
        self.assertEqual(ref_rules_of('速さは要らない。しかし正確さは要る。'), ['neg_then_pos'])
        self.assertEqual(ref_rules_of('まず、電源を、入れて、少し、待ってから、測る。'), ['many_commas'])
        self.assertEqual(ref_rules_of('仕事を休まず。毎日通う。'), ['neg_then_pos'])     # 「休まず」は否定のまま


# ------------------------------------------------------------
# レポートの2回目（2026-10-01。6回目の直しの後、今日の前の版 check_ai_likeness_before.py（md5 cdeccec4）との比べで残っていた物）
#   スキャンの PDF の OCR の崩れの行（教員の再試験・授業の事前演習プリント・本人の手書きノート）の
#   ダッシュ6件を、文字化けと同じ「書き手が打った字ではない」で外した（ocr_noise_near）。崩れの印の無い行も、すぐ前・すぐ後の
#   行が崩れていれば外す。人が打つ半角の記号（ファイル名の _・"…" の引用符1つ・点だけの並び・取り消し線 ~~・Windows のパス）の
#   行の当たりは外れないことも一緒に確かめる
# ------------------------------------------------------------

class Wf7Regressions(unittest.TestCase):
    def assertLevels(self, cases):
        for family, text, want in cases:
            with self.subTest(text=text):
                self.assertEqual(level(family, text), want)

    def test_ocr_noise_lines_are_not_counted(self):
        D, H, Q = chr(0x2014), chr(0x2015), chr(0x22)
        self.assertLevels([
            ('dash', '℃.,.ぁ3。厨の屑~元□府号恥' + H + 'パ・バ', 'miss'),              # 教員の再試験の PDF
            ('dash', '双加' + D + 'Iご！厄ャY/ ハ' + Q * 2 + '11--7-', 'miss'),
            ('dash', D + ' ば' + H + ' さぷ+__', 'miss'),                               # 同じ授業の再試験（期末）
            ('dash', 'ハードウェアはレ訊' + D + ' iJ.tf._」', 'miss'),                 # 事前演習プリント（2）
            ('dash', '．倅屈~.~犀這：）：G---~ 「い' + D + '訂I応象', 'miss'),          # 本人の手書きノート
            # 印の無い行でも、すぐ後の行が崩れていれば崩れの中（再試験の PDF の3行）
            ('dash', '• 印界千\nで(1. :・ 知ク五' + H + '受戎℃ が\n1 ク,~ 瓦孫訊t;.愛殊であも', 'miss'),
            # 絵文字・行末のコロンも同じ原理で外す
            ('emoji', '厨の屑~元 🔺', 'miss'),
            ('colon_end', '屑~元を考える：\nこの経済は閉鎖経済である。', 'miss')])

    def test_human_typed_ascii_marks_keep_hits(self):
        # 人が打つ半角の記号は崩れの印にしない。Claude のデイリーノート・テスト対策 PDF の当たりは当たったまま
        D, H, Q, BS = chr(0x2014), chr(0x2015), chr(0x22), chr(92)
        self.assertLevels([
            ('dash', '・デイリーレポート_20260804_提出用 ' + D + ' Day 4 完成版（Teams提出は本人）', 'hit'),   # ファイル名の _
            ('dash', '補足：外積は「3D」で考える ' + H + ' 向きが' + Q + '別次元' + Q + 'に出る理由', 'hit'),  # 引用符1つずつ
            ('dash', '原因はコンデンサ' + D + '容量の誤差だった。\nいて..の回数を指で数えるのが確実。', 'hit'),  # 点だけの並び
            ('emoji', '・✅ ~~最終発表の手応え・先生からのコメント~~', 'hit'),                     # 取り消し線
            ('emoji', '・D:' + BS + '_退避_20260918' + BS + ' の書類195件\n・📌 繰り越し（変化なし）', 'hit'),  # Windows のパス
            ('emoji', '「えっと...」と言った。\n💡 試験でのポイント', 'hit'),                      # 省略の点
            ('colon_end', '使用した機器を以下に示す：\nオシロスコープ_v1 を使った。', 'hit')])
        # 崩れの印の判定そのもの
        self.assertTrue(cal.ocr_noise('厨の屑~元'))
        self.assertTrue(cal.ocr_noise('さぷ+__'))
        self.assertFalse(cal.ocr_noise('1~2日で終わる'))           # 数字の間の ~ は日本語の字の隣ではない
        self.assertFalse(cal.ocr_noise('40_ツール/_スクリプト/'))   # パスの /_


class Wf2Docx(unittest.TestCase):
    """docx の数式（OMML）と図・OLE の式の跡（〓）の読み取り"""
    M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='report_aik_wf2_')

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def save(self, doc, name):
        path = os.path.join(self.tmp, name)
        doc.save(path)
        return path

    def test_math_paragraph_after_colon(self):
        from docx.oxml import parse_xml
        doc = Document()
        doc.add_paragraph('係数行列 A が正則なら、唯一解は:')
        p = doc.add_paragraph()
        p._p.append(parse_xml('<m:oMathPara xmlns:m="%s"><m:oMath><m:r><m:t>x = A−1b</m:t></m:r></m:oMath></m:oMathPara>'
                              % self.M_NS))
        path = self.save(doc, 'math.docx')
        self.assertIn('x = A−1b', cal.extract_text(path))
        self.assertEqual([f for f in cal.check(path) if f['rule'] == 'colon_end'], [])
        # 式の段落が無ければ拾う
        doc2 = Document()
        doc2.add_paragraph('係数行列 A が正則なら、唯一解は:')
        doc2.add_paragraph('この解は1つだけである。')
        self.assertTrue([f for f in cal.check(self.save(doc2, 'nomath.docx')) if f['rule'] == 'colon_end'])

    def test_ole_object_leaves_a_mark(self):
        # 古い数式エディタの式（w:object）は文字を持たないので〓を置く。読点だけが残って「読点が多い」に当たらない
        from docx.oxml import parse_xml
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('したがって、生産活動は、')
        p._p.append(parse_xml('<w:r xmlns:w="%s"><w:object/></w:r>' % self.W_NS))
        p.add_run('のとき行われ、供給関数は、')
        p._p.append(parse_xml('<w:r xmlns:w="%s"><w:object/></w:r>' % self.W_NS))
        path = self.save(doc, 'ole.docx')
        self.assertEqual(cal.extract_text(path), 'したがって、生産活動は、〓のとき行われ、供給関数は、〓')
        self.assertEqual(cal.check_reference(path), [])


if __name__ == '__main__':
    unittest.main()
