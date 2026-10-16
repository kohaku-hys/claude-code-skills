"""report スキルの機械ゲート（lint / check_ai_likeness / compare_tables / auto_grade_loop /
circuit_lib / run_pipeline）の回帰テスト。2026-09-28 新設。

usage:
    py ~/.claude/skills/report/scripts/tests/test_gates.py        （全部）
    py ~/.claude/skills/report/scripts/tests/test_gates.py -v     （1件ずつ表示）

テスト用の docx / xlsx はその場で python-docx / openpyxl で最小限の物を作り、一時フォルダに置く。
本物のレポートやパターン集（skill-memory）には書き込まない。
"""
import os
import sys
import io
import shutil
import tempfile
import unittest
import subprocess
from decimal import Decimal

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)

from docx import Document                      # noqa: E402
from docx.oxml import parse_xml                 # noqa: E402

import lint                                     # noqa: E402
import check_ai_likeness                        # noqa: E402
import compare_tables                           # noqa: E402

M_DECL = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


# ------------------------------------------------------------
# テスト用の文書づくり
# ------------------------------------------------------------

def add_math(paragraph, inner_xml):
    """段落に Word 数式（OMML）を足す"""
    paragraph._p.append(parse_xml('<m:oMath %s>%s</m:oMath>' % (M_DECL, inner_xml)))


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    return t


def equip_section(doc, header=None):
    """実験方法の章と、使用機器の見出し・4列の表（既定は正しい形）"""
    if header is None:
        header = ['機器・部品', '仕様・型番', '用途', '説明']
    doc.add_paragraph('3. 実験方法')
    doc.add_paragraph('3.0 使用機器')
    doc.add_paragraph('使用した機器を表1に示す。')
    doc.add_paragraph('表1 使用機器')
    add_table(doc, [header, ['直流電源'] + ['-'] * (len(header) - 1)])


def ng_list(issues):
    return [s for s in issues if s.startswith('[NG]')]


class TmpDirCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='report_gates_')

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def path(self, name):
        return os.path.join(self.tmp, name)

    def save(self, doc, name='t.docx'):
        p = self.path(name)
        doc.save(p)
        return p

    def lint_doc(self, doc, **kw):
        issues, stats = lint.lint(self.save(doc), **kw)
        return issues, stats


# ------------------------------------------------------------
# lint.py
# ------------------------------------------------------------

class LintCaptionAndRefs(TmpDirCase):
    def test_caption_with_verb_is_detected(self):
        # 「を測定」を含む正しいキャプションが外れて欠番NGになった記録（skill-memory/report.md）
        doc = Document()
        doc.add_paragraph('図1から図3に結果を示す。')
        doc.add_paragraph('図1 入力波形')
        doc.add_paragraph('図2 位相差を測定するときの入出力波形')
        doc.add_paragraph('図3 出力波形である')
        issues, stats = self.lint_doc(doc)
        self.assertEqual(stats['fig_nums'], [1, 2, 3])
        self.assertFalse([s for s in issues if '飛び番' in s])

    def test_sentence_starting_with_fig_is_not_caption(self):
        doc = Document()
        doc.add_paragraph('図1 回路図')
        doc.add_paragraph('図1 に示す回路で測定した')
        issues, stats = self.lint_doc(doc)
        self.assertEqual(stats['fig_nums'], [1])

    def test_unreferenced_figure_and_table_are_ng(self):
        doc = Document()
        doc.add_paragraph('図1 回路図')
        doc.add_paragraph('表1 測定結果')
        add_table(doc, [['f [Hz]', 'V [V]'], ['100', '1.0']])
        issues, _ = self.lint_doc(doc)
        ng = ng_list(issues)
        self.assertTrue(any('図1 は本文中で参照されていない' in s for s in ng))
        self.assertTrue(any('表1 は本文中で参照されていない' in s for s in ng))

    def test_range_reference_counts_middle_numbers(self):
        doc = Document()
        doc.add_paragraph('結果を図1〜3と表 1 に示す。')
        for n in (1, 2, 3):
            doc.add_paragraph('図%d 波形%d' % (n, n))
        doc.add_paragraph('表 1　測定結果')
        add_table(doc, [['f [Hz]', 'V [V]'], ['100', '1.0']])
        issues, _ = self.lint_doc(doc)
        self.assertFalse([s for s in issues if '参照されていない' in s], issues)


class LintEquations(TmpDirCase):
    def test_numbered_equation_must_be_referenced(self):
        doc = Document()
        p = doc.add_paragraph()
        add_math(p, '<m:r><m:t>V=IR</m:t></m:r>')
        p.add_run('式(1)')
        doc.add_paragraph('　　P ＝ V × I（式2）')
        doc.add_paragraph('式(1)より電流を求めた。')
        issues, stats = self.lint_doc(doc)
        self.assertEqual(stats['eq_nums'], [1, 2])
        ng = ng_list(issues)
        self.assertFalse([s for s in ng if '式(1)' in s])
        self.assertTrue([s for s in ng if '式(2) は本文から参照されていない' in s])

    def test_sentence_and_equation_in_same_paragraph(self):
        # テーマ11: 「…近似的に式1 で表せる。＋改行＋ N = c / n （式1）」が1段落
        doc = Document()
        doc.add_paragraph('点滅回数は近似的に式1 で表せる。\tN = c / n    （式1）')
        issues, stats = self.lint_doc(doc)
        self.assertEqual(stats['eq_nums'], [1])
        self.assertFalse(ng_list(issues))

    def test_reference_without_definition_is_only_warning(self):
        doc = Document()
        doc.add_paragraph('教科書の式(22)より Q を求めた。')
        issues, _ = self.lint_doc(doc)
        self.assertFalse(ng_list(issues))
        self.assertTrue([s for s in issues if s.startswith('[警告]') and '式(22)' in s])


class LintInlineCalc(TmpDirCase):
    """2026-10-07 追加（P-059）: 文の中に埋まった代入計算と、omml_lib の独立行の数式"""

    def calc_ng(self, *paras):
        doc = Document()
        for t in paras:
            doc.add_paragraph(t)
        issues, _ = self.lint_doc(doc)
        return [s for s in ng_list(issues) if 'P-059' in s]

    def test_substitution_in_sentence_is_ng(self):
        # テーマ5 v11 の 4.1・4.2・下書きの形
        self.assertTrue(self.calc_ng('表2の①は R_t = 27.901 ÷ (1 − 27.901 ÷ 3000) ≒ 28.163 Ω となる。'))
        self.assertTrue(self.calc_ng('12 mA のときは 13.4 V ÷ 12 mA ≒ 1120 Ω と小さい。'))
        self.assertTrue(self.calc_ng('境目は √(4.7 × 3000) ≒ 119 Ω。'))
        self.assertTrue(self.calc_ng('差は 4.7 + 0.26 ≒ 5.0 Ω と見積もられる。'))

    def test_ranges_dates_and_symbols_are_not_ng(self):
        self.assertFalse(self.calc_ng(
            '1〜4 mA の R_x は 20.0〜27.5 Ω と小さい。',
            '2026-10-06 に測定した。',
            '(b) と (a) の差は r_A + R_t になる。',
            '例えば R_t = r_V なら ε_a = −50% である。',
            'R = 2.0 / 0.001 = 2000 Ω'))          # 文ではない式だけの行は対象外

    def test_omml_lib_display_and_numbered_lines(self):
        import omml_lib as om
        doc = Document()
        p = doc.add_paragraph('平均値を式(2)に代入して次のように求めた。')
        e = om.cat(om.sub('R', 't'), om.r('='), om.frac('27.901', om.cat(om.r('1−'), om.frac('27.901', '3000'))),
                   om.r('≒'), om.r('28.163'), om.unit('Ω'))
        q = om.insert_equation_after(p, e)
        M = '{%s}' % om.M_NS
        self.assertIsNotNone(q._p.find(M + 'oMathPara'))   # 番号なし＝独立行（分数が小さくならない）
        self.assertIs(q._p.getprevious(), p._p)
        om.add_equation(doc, om.cat(om.sub('R', 't'), om.r('='), om.frac(om.sub('R', 'x'), 'k')), number=2)
        self.assertEqual(doc.element.body[-1].tag.split('}')[1], 'sectPr')   # sectPr は最後のまま
        issues, stats = self.lint_doc(doc)
        self.assertEqual(stats['eq_nums'], [2])             # 番号付きは lint が式として数える
        self.assertFalse([s for s in issues if 'P-051' in s or 'P-059' in s], issues)


class LintSymbols(TmpDirCase):
    def test_sigma_in_body_is_ng(self):
        doc = Document()
        doc.add_paragraph('Σ I = 0 が成り立つ。')
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '「Σ」' in s])

    def test_nary_sum_in_omml_is_ng(self):
        # Word 数式の総和は ∑(U+2211) が m:chr に入る。以前の lint は見ていなかった
        doc = Document()
        p = doc.add_paragraph()
        add_math(p, '<m:nary><m:naryPr><m:chr m:val="∑"/></m:naryPr><m:sub/><m:sup/>'
                    '<m:e><m:r><m:t>I</m:t></m:r></m:e></m:nary>')
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '「Σ」' in s and 'Word数式' in s], issues)

    def test_product_in_omml_text_is_ng(self):
        doc = Document()
        p = doc.add_paragraph()
        add_math(p, '<m:r><m:t>∏ a</m:t></m:r>')
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '「Π」' in s])

    def test_symbol_list_comes_from_range_file(self):
        rf = self.path('range.md')
        with io.open(rf, 'w', encoding='utf-8') as f:
            f.write('## 数学\n### 使用OK\n- Σ（総和）\n### 使用NG（未習）\n- **∮（周回積分）**\n')
        doc = Document()
        doc.add_paragraph('Σ I = 0 と ∮ E dl。')
        issues, _ = self.lint_doc(doc, range_file=rf)
        ng = ng_list(issues)
        self.assertTrue([s for s in ng if '「∮」' in s])
        self.assertFalse([s for s in ng if '「Σ」' in s])

    def test_real_range_file_is_parsed(self):
        symbols, src = lint.load_ng_symbols(lint.RANGE_FILE_DEFAULT)
        for s in ['Σ', 'Π', '∮', '∂', '∇', '∈', '∀', '∃']:
            self.assertIn(s, symbols, src)


class LintTextResidue(TmpDirCase):
    def test_placeholder_is_ng_but_title_brackets_are_not(self):
        doc = Document()
        doc.add_paragraph('【表1の結果について一致度を記述する】')
        doc.add_paragraph('[1] 「【テーマ5】CR 回路の実験書」, 2026.')
        issues, _ = self.lint_doc(doc)
        ng = [s for s in ng_list(issues) if '書きかけ' in s]
        self.assertEqual(len(ng), 1, ng)
        self.assertIn('記述する', ng[0])

    def test_placeholder_in_table_cell(self):
        doc = Document()
        doc.add_paragraph('表1 結果')
        add_table(doc, [['V [V]', '備考'], ['1.0', '（要記入）']])
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '書きかけ' in s])

    def test_unicode_super_and_subscripts_are_ng(self):
        doc = Document()
        doc.add_paragraph('電流は 1.0×10⁻³ A、入力は V₁ とする。')
        issues, _ = self.lint_doc(doc)
        ng = [s for s in ng_list(issues) if 'Unicode' in s]
        self.assertEqual(len(ng), 1)
        self.assertIn('3 個', ng[0])   # ⁻ ³ ₁

    def test_plain_underscore_notation_is_ng(self):
        doc = Document()
        doc.add_paragraph('電圧 V_FB を測った。10^{-3} A。')
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '文字のまま' in s and '2 箇所' in s])

    def test_code_listing_is_not_checked(self):
        doc = Document()
        doc.add_paragraph('リスト1 main.c')
        add_table(doc, [['int in_num; x_1 = 0; /* TODO */ ;;']])
        issues, _ = self.lint_doc(doc)
        self.assertFalse([s for s in ng_list(issues) if '書きかけ' in s or '文字のまま' in s])

    def test_omml_subscript_needs_subscript_run(self):
        doc = Document()
        p = doc.add_paragraph()
        add_math(p, '<m:sSub><m:e><m:r><m:t>V</m:t></m:r></m:e><m:sub><m:r><m:t>out</m:t></m:r></m:sub></m:sSub>')
        doc.add_paragraph('出力電圧 Vout を測った。')
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '下付き書式の文字が1つも無い' in s])
        # 本文に下付きの run があれば通る
        q = doc.add_paragraph('出力電圧 V')
        r = q.add_run('out')
        r.font.subscript = True
        issues, _ = self.lint_doc(doc)
        self.assertFalse([s for s in ng_list(issues) if '下付き書式の文字が1つも無い' in s])


class LintEquipment(TmpDirCase):
    def test_correct_equipment_table_passes(self):
        doc = Document()
        equip_section(doc)
        issues, _ = self.lint_doc(doc)
        self.assertFalse(ng_list(issues), issues)

    def test_three_column_table_is_ng(self):
        doc = Document()
        equip_section(doc, ['機器・部品', '仕様', '説明'])
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '使用機器の表の列' in s])

    def test_missing_heading_is_ng(self):
        doc = Document()
        doc.add_paragraph('3. 実験方法')
        doc.add_paragraph('表1 使用機器')
        add_table(doc, [['機器・部品', '仕様・型番', '用途', '説明']])
        doc.add_paragraph('機器を表1に示す。')
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '「使用機器」の見出しが無い' in s])

    def test_heading_without_table_is_ng(self):
        doc = Document()
        doc.add_paragraph('3. 実験方法')
        doc.add_paragraph('3.3 使用機器・部品')
        doc.add_paragraph('・直流電源')
        doc.add_paragraph('3.4 実験1')
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if '見出しの後に表が無い' in s])

    def test_no_method_chapter_skips_check(self):
        doc = Document()
        doc.add_paragraph('1. 概要')
        issues, stats = self.lint_doc(doc)
        self.assertFalse(ng_list(issues))
        self.assertTrue([s for s in stats['infos'] if '省略' in s])


class LintDocStructure(TmpDirCase):
    """Word では正しく見えるが構造がおかしい物（2026-09-29 学生実験のレポートの最終監査で発見・P-051/P-052）"""

    SUB = '<m:sSub><m:e><m:r><m:t>R</m:t></m:r></m:e><m:sub><m:r><m:t>t</m:t></m:r></m:sub></m:sSub>'

    def test_element_nested_in_math_text_is_ng(self):
        # sSub の断片を r() で包んで sSup の底にした（m:t の中に m:sSub が入る）
        doc = Document()
        p = doc.add_paragraph()
        add_math(p, '<m:sSup><m:e><m:r><m:t>%s</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>' % self.SUB)
        issues, _ = self.lint_doc(doc)
        self.assertTrue([s for s in ng_list(issues) if 'm:t の中' in s], issues)

    def test_well_formed_math_is_not_flagged(self):
        doc = Document()
        p = doc.add_paragraph()
        add_math(p, '<m:sSup><m:e>%s</m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>' % self.SUB)
        issues, _ = self.lint_doc(doc)
        self.assertFalse([s for s in issues if 'm:t の中' in s], issues)

    def test_default_core_properties_warn_and_setting_them_clears(self):
        import datetime
        doc = Document()                      # python-docx の雛形は作成者 python-docx・作成日 2013-12-23
        issues, _ = self.lint_doc(doc)
        warn = [s for s in issues if s.startswith('[警告]') and 'python-docx' in s]
        self.assertTrue(warn, issues)
        self.assertFalse([s for s in ng_list(issues) if 'python-docx' in s])   # 警告であって NG ではない
        doc.core_properties.author = '山田 太郎'
        doc.core_properties.created = datetime.datetime(2026, 9, 29, 12, 0, 0)
        issues, _ = self.lint_doc(doc)
        self.assertFalse([s for s in issues if 'python-docx' in s], issues)


class PdfLayoutCheck(unittest.TestCase):
    """PDF の折り返しの検査（2026-09-29 最終監査。実際に出た折り返しをそのまま使う）"""

    def test_long_vowel_at_line_head_is_flagged(self):
        import pdf_layout_check as pl
        self.assertTrue(pl.bad_breaks('…ホスト名は Web サーバ', 'ーには PHP 内蔵の簡易サーバーを使用し'))
        self.assertTrue(pl.bad_breaks('V_BE はデジタルマルチメ', 'ータで測定し'))

    def test_iv_split_and_url_split_are_flagged(self):
        import pdf_layout_check as pl
        self.assertTrue(pl.bad_breaks('(a) と (b) の I', '－V 特性は同じグラフに描く。'))
        self.assertTrue(pl.bad_breaks('PC のブラウザで http://pi-', 'elec-00x0000.local/ファイル名'))

    def test_normal_breaks_are_not_flagged(self):
        import pdf_layout_check as pl
        self.assertFalse(pl.bad_breaks('実験は表3の順に進める。Web ページのファイルは', 'すべてホームディレクトリの theme8 に置き、'))
        self.assertFalse(pl.bad_breaks('図1の回路で、', '電流を変える。'))
        self.assertFalse(pl.bad_breaks('', 'ーの前が空の行'))


class LintExitCode(TmpDirCase):
    def test_cli_exit_1_on_ng_and_0_when_clean(self):
        bad = Document()
        bad.add_paragraph('図1 回路図')
        good = Document()
        good.add_paragraph('回路を図1に示す。')
        good.add_paragraph('図1 回路図')
        for doc, name, want in ((bad, 'bad.docx', 1), (good, 'good.docx', 0)):
            p = self.save(doc, name)
            r = subprocess.run([sys.executable, os.path.join(SCRIPTS, 'lint.py'), p],
                               capture_output=True, text=True, encoding='utf-8', errors='replace')
            self.assertEqual(r.returncode, want, r.stdout)


# ------------------------------------------------------------
# check_ai_likeness.py
# ------------------------------------------------------------

class AiLikeness(TmpDirCase):
    def findings(self, text):
        doc = Document()
        doc.add_paragraph(text)
        return check_ai_likeness.check(self.save(doc))

    def test_niokite_plain_and_mo_counted_once_each(self):
        f = self.findings('回路において電流を測る。回路においても同じ。')
        matches = [x['match'] for x in f]
        self.assertEqual(matches.count('において'), 1)
        self.assertEqual(matches.count('においても'), 1)

    def test_shared_doc_rules_phrases(self):
        for text in ['社会に貢献する。', 'さらなる検討が必要である。', '確認が重要である。',
                     '精度が求められる。', '妥当と言えるだろう。', '妥当といえるだろう。']:
            f = self.findings(text)
            self.assertTrue([x for x in f if x['pid'] == 'SDR'], text)

    def test_koushouchi_is_flagged(self):
        # 本人が使わない語（2026-10-13）
        f = self.findings('R_C の実際の値が公称値の 3.6 kΩ より大きい。')
        self.assertTrue([x for x in f if x['match'] == '公称値'])

    def test_physics_sarani_is_not_flagged(self):
        f = self.findings('電位障壁がさらなる拡散を妨げる。')
        self.assertFalse([x for x in f if x['pid'] == 'SDR'])

    def test_kadai_heading_quote_is_not_flagged(self):
        # 問題文どおりの見出しは書き手の文ではない。本文の同じ言い回しは従来どおり止める
        doc = Document()
        doc.add_paragraph('5.7 検討課題7　これはスイッチ on の状態であることを示している。')
        doc.add_paragraph('V_CE はほぼ 0 である。')
        f = check_ai_likeness.check(self.save(doc))
        self.assertFalse([x for x in f if x['pid'] == 'P-025'])
        f = self.findings('この結果はスイッチ on であることを示している。')
        self.assertTrue([x for x in f if x['pid'] == 'P-025'])

    def test_paren_citation_density(self):
        # P-066: 括弧で添えた（表N）（N.N節）が多いと止める。文の中の参照は数えない
        many = '電流は流れなかった（表3～表5）。上限に近い（5.4節）。破線の下に並んだ（図6）。値は 0.14 V（表6）である。'
        self.assertTrue([x for x in self.findings(many) if x['pid'] == 'P-066'])
        few = '表6で V_CE が 15 V になる。4.3節では R_C が 3.63 kΩ と出ている。図6に◆で書き込んだ。値は 1 V であった（表3）。'
        self.assertFalse([x for x in self.findings(few) if x['pid'] == 'P-066'])


# ------------------------------------------------------------
# compare_tables.py
# ------------------------------------------------------------

class NumberParsing(unittest.TestCase):
    def vals(self, text):
        return [n.value for n in compare_tables.parse_numbers(text)]

    def test_basic_forms(self):
        self.assertEqual(self.vals('1,000'), [Decimal('1000')])
        self.assertEqual(self.vals('−5'), [Decimal('-5')])
        self.assertEqual(self.vals('0-10'), [Decimal('0'), Decimal('10')])
        self.assertEqual(self.vals('R1 と CH2'), [])
        self.assertNotIn(Decimal('1815'), self.vals('2SC1815'))
        self.assertEqual(self.vals('１．２３'), [Decimal('1.23')])

    def test_powers_of_ten(self):
        self.assertEqual(self.vals('1.2×10^(-3)'), [Decimal('0.0012')])
        self.assertEqual(self.vals('10⁻³'), [Decimal('0.001')])
        self.assertEqual(self.vals('3.3e-3'), [Decimal('0.0033')])

    def test_subscripts_are_not_numbers(self):
        self.assertEqual(self.vals('V₁ = 3.2'), [Decimal('3.2')])
        self.assertEqual(self.vals('V_(1) = 3.2'), [Decimal('3.2')])

    def test_decimals_are_kept(self):
        n = compare_tables.parse_numbers('0.100')[0]
        self.assertEqual(n.value.as_tuple().exponent, -3)


def make_xlsx(path, rows, fmt=None, sheet='data'):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            if v is None:
                continue
            cell = ws.cell(row=r, column=c, value=v)
            if fmt and isinstance(v, (int, float)) and r > 1:
                cell.number_format = fmt
    wb.save(path)


def make_report(path, rows, caption='表2 測定結果', with_equipment=True):
    doc = Document()
    if with_equipment:
        doc.add_paragraph('表1 使用機器')
        add_table(doc, [['機器・部品', '仕様・型番', '用途', '説明'], ['抵抗', '1 kΩ', '負荷', '-']])
    doc.add_paragraph(caption)
    add_table(doc, rows)
    doc.save(path)


SRC_ROWS = [['f [Hz]', 'V [V]', 'I [mA]'],
            [100, 1.23, 0.5],
            [200, 2.34, 0.61],
            [500, 3.45, 0.72],
            [1000, 4.56, None],
            [2000, 5.67, 0.94]]


def report_rows(changes=None):
    rows = [['f [Hz]', 'V [V]', 'I [mA]'],
            ['100', '1.23', '0.5'],
            ['200', '2.34', '0.61'],
            ['500', '3.45', '0.72'],
            ['1000', '4.56', ''],
            ['2000', '5.67', '0.94']]
    for (r, c), v in (changes or {}).items():
        rows[r][c] = v
    return rows


class CompareTables(TmpDirCase):
    def run_compare(self, rows, **kw):
        src = self.path('src.xlsx')
        make_xlsx(src, SRC_ROWS)
        out = self.path('out.docx')
        make_report(out, rows)
        return compare_tables.compare(out, [src], **kw)

    def test_exact_copy_passes(self):
        ng, warn, info, notes = self.run_compare(report_rows())
        self.assertEqual(ng, [])
        self.assertTrue([s for s in notes if '除外: 表1（使用機器の表）' in s])

    def test_fabricated_value_is_ng(self):
        ng, _w, _i, _n = self.run_compare(report_rows({(2, 1): '9.99'}))
        self.assertEqual(len(ng), 1)
        self.assertIn('9.99', ng[0])

    def test_value_moved_from_elsewhere_is_ng(self):
        # 2.34 は入力の別の場所にある値。列の対応から「値の食い違い」として捕まえる
        ng, _w, _i, _n = self.run_compare(report_rows({(1, 1): '2.34'}))
        self.assertEqual(len(ng), 1)
        self.assertIn('値の食い違い', ng[0])

    def test_filled_blank_is_ng(self):
        # 入力で空欄の I(1000 Hz) に、入力の別の場所にある値 0.5 を入れた
        ng, _w, _i, _n = self.run_compare(report_rows({(4, 2): '0.5'}))
        self.assertEqual(len(ng), 1)
        self.assertIn('空欄', ng[0])

    def test_rounding_typed_value_is_ng(self):
        ng, _w, _i, _n = self.run_compare(report_rows({(3, 1): '3.5'}))
        self.assertEqual(len(ng), 1)
        self.assertIn('丸め', ng[0])

    def test_skip_table_and_skip_col(self):
        ng, _w, _i, notes = self.run_compare(report_rows({(2, 1): '9.99'}), skip_tables=[2])
        self.assertEqual(ng, [])
        self.assertTrue([s for s in notes if '表2（--skip-table' in s])
        ng, _w, _i, _n = self.run_compare(report_rows({(2, 1): '9.99'}), skip_cols={2: {2}})
        self.assertEqual(ng, [])

    def test_general_format_trailing_zero_is_ok(self):
        # Excel の General は 1.30 を 1.3 でしか保存できないので、1.30 と書いても桁数違いにしない
        ng, warn, _i, _n = self.run_compare(report_rows({(1, 2): '0.50'}))
        self.assertEqual(ng, [])
        self.assertEqual(warn, [])

    def test_formatted_cell_trailing_zero_is_warning(self):
        # 入力は 0.5 と打って表示書式 0.000（=0.500）。出力の 0.50 は打った値とも表示とも桁数が違う
        src = self.path('src.xlsx')
        make_xlsx(src, SRC_ROWS, fmt='0.000')
        out = self.path('out.docx')
        make_report(out, report_rows({(1, 2): '0.50'}))
        ng, warn, _i, _n = compare_tables.compare(out, [src])
        self.assertEqual(ng, [])
        self.assertTrue([s for s in warn if '末尾の0' in s and '0.50' in s], warn)
        # 表示どおり（0.500）か打った値どおり（0.5）なら警告なし
        for text in ('0.500', '0.5'):
            make_report(out, report_rows({(1, 2): text}))
            ng, warn, _i, _n = compare_tables.compare(out, [src])
            self.assertEqual((ng, warn), ([], []), text)

    def test_formula_rounding_is_info(self):
        sc = compare_tables.SrcCell('x!A1', [Decimal('255.55555555555554')], True, '255.5')
        idx = compare_tables.SourceIndex([sc])
        kind, _sc = idx.best_match(compare_tables.parse_numbers('256')[0])
        self.assertEqual(kind, compare_tables.ROUND_F)
        sc2 = compare_tables.SrcCell('x!A2', [Decimal('8.32')], False, '8.32')
        idx2 = compare_tables.SourceIndex([sc2])
        kind2, _ = idx2.best_match(compare_tables.parse_numbers('8.3')[0])
        self.assertEqual(kind2, compare_tables.ROUND_C)

    def test_unit_conversion_is_warning(self):
        sc = compare_tables.SrcCell('x!A1', [Decimal('0.0012')], False, '0.0012')
        idx = compare_tables.SourceIndex([sc])
        kind, _sc = idx.best_match(compare_tables.parse_numbers('1.2')[0])
        self.assertEqual(kind, compare_tables.SCALED)

    def test_docx_source(self):
        src = self.path('pre.docx')
        make_report(src, report_rows(), with_equipment=False)
        out = self.path('out.docx')
        make_report(out, report_rows())
        ng, warn, _i, _n = compare_tables.compare(out, [src])
        self.assertEqual(ng, [])
        make_report(out, report_rows({(2, 1): '2.3'}))
        ng, _w, _i, _n = compare_tables.compare(out, [src])
        self.assertEqual(len(ng), 1)

    def test_cli_exit_codes(self):
        src = self.path('src.xlsx')
        make_xlsx(src, SRC_ROWS)
        out = self.path('out.docx')
        make_report(out, report_rows({(2, 1): '9.99'}))
        cmd = [sys.executable, os.path.join(SCRIPTS, 'compare_tables.py'), out, '--source', src]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace')
        self.assertEqual(r.returncode, 1, r.stdout)
        r = subprocess.run(cmd + ['--skip-table', '2'], capture_output=True, text=True,
                           encoding='utf-8', errors='replace')
        self.assertEqual(r.returncode, 0, r.stdout)
        r = subprocess.run(cmd[:-1] + [self.path('none.xlsx')], capture_output=True, text=True,
                           encoding='utf-8', errors='replace')
        self.assertEqual(r.returncode, 2, r.stdout)


# ------------------------------------------------------------
# auto_grade_loop.py
# ------------------------------------------------------------

GRADE_MD = """# 採点結果

総合評価：**{score}/100点**

### 1章 目的（20点満点）→ 18/20
**減点理由（−2点）**：目的が長い
"""


class AutoGradeLoop(TmpDirCase):
    def run_extract(self, score):
        md = self.path('grade.md')
        with io.open(md, 'w', encoding='utf-8') as f:
            f.write(GRADE_MD.format(score=score))
        patterns = self.path('patterns.md')
        with io.open(patterns, 'w', encoding='utf-8') as f:
            f.write('# パターン\n')
        r = subprocess.run([sys.executable, os.path.join(SCRIPTS, 'auto_grade_loop.py'), 'extract', md,
                            '--theme', 'テスト', '--patterns', patterns],
                           capture_output=True, text=True, encoding='utf-8', errors='replace')
        with io.open(patterns, encoding='utf-8') as f:
            self.assertEqual(f.read(), '# パターン\n')   # --save なしでは書かない
        return r

    def test_below_90_exits_3(self):
        r = self.run_extract(85)
        self.assertEqual(r.returncode, 3, r.stdout + r.stderr)

    def test_90_or_more_exits_0(self):
        r = self.run_extract(92)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)


# ------------------------------------------------------------
# circuit_lib.py
# ------------------------------------------------------------

class CircuitLib(TmpDirCase):
    def test_save_circuit_refuses_overlapping_labels(self):
        import circuit_lib
        fig, ax = circuit_lib.new_circuit(figsize=(4, 3))
        ax.set_xlim(0, 4)
        ax.set_ylim(0, 3)
        circuit_lib.draw_wire(ax, [(0, 1.5), (4, 1.5)])
        circuit_lib.draw_label(ax, 2, 1.5, 'R = 10 kΩ')
        out = self.path('bad.png')
        with self.assertRaises(AssertionError):
            circuit_lib.save_circuit(fig, out)
        self.assertFalse(os.path.exists(out))

    def test_save_circuit_writes_clean_figure(self):
        import circuit_lib
        fig, ax = circuit_lib.new_circuit(figsize=(4, 3))
        ax.set_xlim(0, 4)
        ax.set_ylim(0, 3)
        circuit_lib.draw_wire(ax, [(0, 1), (4, 1)])
        circuit_lib.draw_label(ax, 2, 2.2, 'R = 10 kΩ')
        circuit_lib.assert_no_overlaps(fig, ax)   # 先に呼んでいても二重呼び出しで壊れない
        out = self.path('ok.png')
        circuit_lib.save_circuit(fig, out)
        self.assertTrue(os.path.exists(out))


# ------------------------------------------------------------
# run_pipeline.py
# ------------------------------------------------------------

class Pipeline(TmpDirCase):
    def test_source_runs_compare_and_fails_on_fabrication(self):
        src = self.path('src.xlsx')
        make_xlsx(src, SRC_ROWS)
        out = self.path('out.docx')
        make_report(out, report_rows({(2, 1): '9.99'}))
        r = subprocess.run([sys.executable, os.path.join(SCRIPTS, 'run_pipeline.py'), out, '--source', src],
                           capture_output=True, text=True, encoding='utf-8', errors='replace')
        self.assertEqual(r.returncode, 1)
        with io.open(self.path('out_pipeline_report.md'), encoding='utf-8') as f:
            rep = f.read()
        self.assertIn('STEP 4b-4 表データ照合 | **NG**', rep)

    def test_without_source_records_not_run(self):
        out = self.path('out.docx')
        make_report(out, report_rows())
        subprocess.run([sys.executable, os.path.join(SCRIPTS, 'run_pipeline.py'), out],
                       capture_output=True, text=True, encoding='utf-8', errors='replace')
        with io.open(self.path('out_pipeline_report.md'), encoding='utf-8') as f:
            self.assertIn('表データ照合は未実行', f.read())

    def test_cutout_warning(self):
        import fitz
        import run_pipeline
        pdf = self.path('youryou.pdf')
        d = fitz.open()
        page = d.new_page()
        page.insert_text((50, 80), '教科書の図を切り取って貼り付けることは禁止する。', fontname='japan', fontsize=11)
        d.save(pdf)
        # 名前に pdf を含む画像（extract_pdf_figures.py の出力名の形）を貼った docx
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        png = self.path('pdf_fig3_circuit.png')
        fig = plt.figure(figsize=(1, 1))
        fig.savefig(png)
        plt.close(fig)
        doc = Document()
        doc.add_picture(png)
        out = self.save(doc, 'fig.docx')
        msg = run_pipeline.check_cutout(out, pdf)
        self.assertIn('切り出し禁止らしい文がある', msg)
        self.assertIn('pdf_fig3_circuit.png', msg)


# ------------------------------------------------------------
# 保存名（SKILL.md STEP 5 が使う shared-doc-rules の safe_save_path。読むだけで直さない）
# ------------------------------------------------------------

class SaveName(TmpDirCase):
    def test_report_names_follow_version_rule(self):
        sdr = os.path.join(os.path.dirname(os.path.dirname(SCRIPTS)), 'shared-doc-rules', 'scripts')
        if not os.path.exists(os.path.join(sdr, 'safe_save_path.py')):
            self.skipTest('shared-doc-rules/scripts/safe_save_path.py が無い')
        sys.path.insert(0, sdr)
        import safe_save_path
        if not hasattr(safe_save_path, 'check_file_name'):
            self.skipTest('check_file_name が無い版の safe_save_path')
        good = self.path('テーマ10_本レポート_v6.docx')
        self.assertEqual(safe_save_path.safe_save_path(good), good)
        for bad in ('テーマ10_本レポート_最終版_v6.docx', 'テーマ10_本レポート.docx'):
            with self.assertRaises(ValueError):
                safe_save_path.safe_save_path(self.path(bad))



class TestRichTextAndLayout(unittest.TestCase):
    """2026-09-29 追加: 識別子の下付き誤変換と、図表とキャプションの泣き別れ"""

    def test_identifier_not_subscript(self):
        import rich_text
        d = Document()
        p = d.add_paragraph()
        rich_text.add_rich(p, 'read_adc() と SPI_MOSI、RPi_CAE')
        self.assertFalse(any(r.font.subscript for r in p.runs))

    def test_physics_subscript_kept(self):
        import rich_text
        d = Document()
        p = d.add_paragraph()
        rich_text.add_rich(p, 'V_{BE} と R_{C}I_{C}、V_a')
        subs = [r.text for r in p.runs if r.font.subscript]
        self.assertEqual(subs, ['BE', 'C', 'C', 'a'])

    def test_figure_and_table_caption_keep_together(self):
        import table_lib
        from PIL import Image
        tmp = tempfile.mkdtemp()
        try:
            png = os.path.join(tmp, 'x.png')
            Image.new('RGB', (40, 20), 'white').save(png)
            d = Document()
            table_lib.add_figure(d, png, '図1　テスト', width_cm=3)
            pic_par = d.paragraphs[-2]
            self.assertTrue(pic_par.paragraph_format.keep_with_next)
            cap = table_lib.add_table_caption(d, '表1　テスト')
            self.assertTrue(cap.paragraph_format.keep_with_next)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

if __name__ == '__main__':
    unittest.main()
