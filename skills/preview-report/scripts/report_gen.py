"""
予習レポート Word (.docx) 生成ライブラリ

動画解析結果 (dict) から標準構成の Word を組み立てる。
omml_lib を使って OMML 数式を正しく描画。

使い方:
    from report_gen import build_report, ReportConfig, Section, MathBlock

    report = ReportConfig(
        subject="線形代数Ⅰ",
        subtitle="第5回：§2.2 行列式の性質と一般の行列式の計算",
        videos=[
            {"id": "05_1", "title": "行列式の性質（スカラー倍・次数下げ）", "duration": "19:10"},
            {"id": "05_2", "title": "行列式の性質（入れ替え・定数倍）", "duration": "15:27"},
            {"id": "05_3", "title": "一般の行列式の計算", "duration": "11:16"},
        ],
        goals=[
            "§2.2 の7つの定理（2.2.1〜2.2.6 + 系2.2.7）を体系的に理解する",
            "…",
        ],
        sections=[
            Section(heading="§1 導入", paragraphs=["…"]),
            Section(heading="§2 スカラー倍", paragraphs=["…"],
                    math=[MathBlock(caption="定理2.2.1", inner_xml=m_det(...))])
        ],
        problems=[...],
        summary_paragraphs=["…"],
    )
    path = build_report(report, output_dir="…/予習/", filename_base="線形代数Ⅰ_予習レポート_05")
"""
import os, datetime
from pathlib import Path
from dataclasses import dataclass, field
from typing import List, Optional, Union
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

# このモジュールのディレクトリを import パスへ
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from omml_lib import (
    setup_japanese, heading, add_math, safe_save_path,
    m_r, m_ri  # 最小再エクスポート（呼び出し側が import 省略できる）
)


# === データクラス ===

@dataclass
class MathBlock:
    """数式ブロック。caption は式の上に書く説明（省略可）"""
    inner_xml: str                   # omml_lib の m_* で組んだ XML 文字列
    caption: Optional[str] = None    # 式の前の説明文
    center: bool = True

@dataclass
class Section:
    """§1, §2, ... のセクション 1 つ"""
    heading: str                              # 例: "§2 行列式の性質（動画 05_1）"
    paragraphs: List[str] = field(default_factory=list)
    math: List[MathBlock] = field(default_factory=list)
    subsections: List["Section"] = field(default_factory=list)
    bullets: List[str] = field(default_factory=list)

@dataclass
class Problem:
    """練習問題 1 題"""
    question: str                     # 問題文
    math: Optional[str] = None        # 数式部分 (OMML XML)
    answer: str = ""                  # 解答・解説

@dataclass
class ReportConfig:
    """予習レポート全体の構成"""
    subject: str
    subtitle: str
    videos: List[dict] = field(default_factory=list)     # [{id, title, duration}]
    goals: List[str] = field(default_factory=list)
    sections: List[Section] = field(default_factory=list)
    common_pitfalls: List[str] = field(default_factory=list)
    problems: List[Problem] = field(default_factory=list)
    summary_paragraphs: List[str] = field(default_factory=list)
    # メタ情報（表紙に入る）。ReportConfig に直接渡すか、環境変数
    # PREVIEW_REPORT_STUDENT_ID / PREVIEW_REPORT_STUDENT_NAME / PREVIEW_REPORT_AFFILIATION で既定値を決める。
    # どちらも無ければダミー値が入る
    student_id: str = field(default_factory=lambda: os.environ.get("PREVIEW_REPORT_STUDENT_ID", "s00x0000"))
    student_name: str = field(default_factory=lambda: os.environ.get("PREVIEW_REPORT_STUDENT_NAME", "山田 太郎"))
    affiliation: str = field(default_factory=lambda: os.environ.get("PREVIEW_REPORT_AFFILIATION", "〇〇学部 〇〇学科"))
    # 動画キャプチャの manifest.json（preview.py capture / capture_folder.py が作る）のパス。
    # 無いと「動画を見ずに作った」疑いとして警告する（skill-memory: タイトルが明確でも動画を必ず見る）
    capture_manifest: Optional[str] = None


# === 標準構成の検査（SKILL.md「レポート生成の原則」の標準構成） ===

PROBLEMS_MIN, PROBLEMS_MAX = 5, 6

# 学習範囲のホワイトリスト。読むだけで編集しない（正本は skill-memory 側。report 側に置かれた場合も読む）
_HOME = os.environ.get("USERPROFILE") or str(Path.home())
SCOPE_FILES = [
    os.path.join(_HOME, ".claude", "skill-memory", "学習範囲.md"),
    os.path.join(_HOME, ".claude", "skills", "report", "references", "学習範囲.md"),
]
# OMML の演算子は Σ/Π ではなく ∑/∏ で入るので、同じ記号として扱う
SYMBOL_ALIASES = {"Σ": "Σ∑", "Π": "Π∏"}


def load_ng_symbols():
    """学習範囲.md の「使用NG」節から、太字で書かれた1文字の記号を集める。
    ファイルが無ければ空のリスト（検査しない）と、読んだファイルのパスを返す。"""
    import re
    for path in SCOPE_FILES:
        if not os.path.exists(path):
            continue
        with open(path, encoding="utf-8") as f:
            text = f.read()
        m = re.search(r"###\s*使用NG[^\n]*\n(.*?)(?=\n###|\n---|\Z)", text, re.S)
        if not m:
            return [], path
        symbols = []
        for bold in re.findall(r"\*\*([^*]+)\*\*", m.group(1)):
            head = re.split(r"[（(]", bold)[0]      # 「∈, ∋（要素記号）」→「∈, ∋」
            for tok in re.split(r"[,、\s]+", head):
                if len(tok) == 1:                  # 「ベクトル記号」「行列・行列式」は記号ではないので外す
                    symbols.append(tok)
        return symbols, path
    return [], None


def _all_text(cfg: "ReportConfig"):
    """記号の検査に使う全テキスト（本文・箇条書き・数式 XML・問題・まとめ）。"""
    parts = [cfg.subtitle] + list(cfg.goals) + list(cfg.common_pitfalls) + list(cfg.summary_paragraphs)

    def walk(sec):
        parts.append(sec.heading)
        parts.extend(sec.paragraphs)
        parts.extend(sec.bullets)
        for mb in sec.math:
            parts.append(mb.caption or "")
            parts.append(mb.inner_xml)
        for sub in sec.subsections:
            walk(sub)

    for sec in cfg.sections:
        walk(sec)
    for q in cfg.problems:
        parts += [q.question, q.math or "", q.answer]
    return "\n".join(parts)


def _warn(msg: str):
    """警告を stderr に UTF-8 で出す。パイプ先の既定は cp932 なので、print だと日本語が化ける。"""
    data = f"[preview-report 警告] {msg}\n".encode("utf-8")
    try:
        sys.stderr.flush()
        sys.stderr.buffer.write(data)
        sys.stderr.buffer.flush()
    except AttributeError:   # buffer を持たない stderr（IDE 等）
        print(f"[preview-report 警告] {msg}", file=sys.stderr)


def check_config(cfg: "ReportConfig"):
    """標準構成を検査して (止める理由のリスト, 警告のリスト) を返す。

    止める: 学習目標・本文の節・つまずきやすい点・練習問題（解答付き5〜6問）・授業前のまとめ が欠けている
    警告:   動画キャプチャの manifest が無い／学習範囲の外の記号を使っている
    """
    import json
    errors, warnings = [], []
    if not cfg.goals:
        errors.append("■学習目標 が無い（goals）")
    if not cfg.sections:
        errors.append("§1 導入〜各動画の内容 の節が無い（sections）")
    if not cfg.common_pitfalls:
        errors.append("■つまずきやすい点 が無い（common_pitfalls）")
    n = len(cfg.problems)
    if n < PROBLEMS_MIN or n > PROBLEMS_MAX:
        errors.append(f"練習問題が {n} 問（標準構成は解答付き {PROBLEMS_MIN}〜{PROBLEMS_MAX} 問）")
    no_answer = [str(i) for i, q in enumerate(cfg.problems, 1) if not q.answer.strip()]
    if no_answer:
        errors.append(f"解答の無い練習問題がある（問{', 問'.join(no_answer)}）")
    if not cfg.summary_paragraphs:
        errors.append("■授業前のまとめ が無い（summary_paragraphs）")

    # 動画を見たか（manifest の有無で代理判定）
    if not cfg.capture_manifest:
        warnings.append("capture_manifest が未指定。タイトルが明確でも動画を必ず見る（skill-memory）。"
                        "preview.py capture / capture_folder.py の manifest.json を渡す")
    elif not os.path.exists(cfg.capture_manifest):
        warnings.append(f"capture_manifest が見つからない: {cfg.capture_manifest}")
    else:
        try:
            with open(cfg.capture_manifest, encoding="utf-8") as f:
                json.load(f)
        except (OSError, ValueError) as e:
            warnings.append(f"capture_manifest が JSON として読めない: {e}")

    # 学習範囲の外の記号（線形代数など、その記号自体が授業の中身のこともあるので止めずに警告）
    ng_symbols, scope_path = load_ng_symbols()
    if ng_symbols:
        text = _all_text(cfg)
        used = []
        for s in ng_symbols:
            forms = SYMBOL_ALIASES.get(s, s)
            if any(ch in text for ch in forms):
                used.append(s)
        if used:
            warnings.append(f"学習範囲の外の記号 {' '.join(used)} を使っている（{os.path.basename(scope_path)}）。"
                            "その回の授業で習う記号でなければ展開形に書き換える")
    return errors, warnings


# === Word 組み立て ===

def _title_block(doc, cfg: ReportConfig):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{cfg.subject} 予習レポート")
    r.bold = True; r.font.size = Pt(16); r.font.name = 'Yu Gothic'
    rpr = r._element.get_or_add_rPr()
    rF = OxmlElement('w:rFonts'); rF.set(qn('w:eastAsia'), 'Yu Gothic'); rpr.append(rF)

    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(cfg.subtitle)
    r2.bold = True; r2.font.size = Pt(13); r2.font.name = 'Yu Gothic'
    rpr2 = r2._element.get_or_add_rPr()
    rF2 = OxmlElement('w:rFonts'); rF2.set(qn('w:eastAsia'), 'Yu Gothic'); rpr2.append(rF2)

    doc.add_paragraph()
    today = datetime.date.today().isoformat()
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p3.add_run(f"学籍番号: {cfg.student_id}　氏名: {cfg.student_name}\n"
               f"提出日: {today}\n{cfg.affiliation}")
    doc.add_paragraph()


def _bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.add_run(text)
    return p


def _write_section(doc, sec: Section, level: int = 0):
    size = 12 if level == 0 else 11
    heading(doc, sec.heading, size=size)
    for para in sec.paragraphs:
        doc.add_paragraph(para)
    for bl in sec.bullets:
        _bullet(doc, bl)
    for mb in sec.math:
        if mb.caption:
            doc.add_paragraph(mb.caption)
        add_math(doc, mb.inner_xml, center=mb.center)
    for sub in sec.subsections:
        _write_section(doc, sub, level + 1)


def build_report(cfg: ReportConfig, output_dir: str, filename_base: str) -> str:
    """
    レポート Word を生成して保存パスを返す。
    filename_base: 拡張子なし。最終的に `<base>_v1.docx` として保存（OneDrive ロック時は v2+）。
    標準構成が欠けていたら、Word を作る前に ValueError で止める（check_config）。
    """
    errors, warnings = check_config(cfg)
    for w in warnings:
        _warn(w)
    if errors:
        raise ValueError("予習レポートの標準構成が欠けているので作らない: " + " / ".join(errors))

    doc = Document()
    setup_japanese(doc)

    _title_block(doc, cfg)

    # 対応動画
    if cfg.videos:
        heading(doc, "■ 対応動画（反転授業用）")
        for v in cfg.videos:
            line = f"{v.get('id','')}: {v.get('title','')}"
            if v.get('duration'):
                line += f" [{v['duration']}]"
            _bullet(doc, line)

    # 学習目標
    if cfg.goals:
        heading(doc, "■ 学習目標")
        for g in cfg.goals:
            _bullet(doc, g)

    # 本文セクション
    for sec in cfg.sections:
        _write_section(doc, sec)

    # つまずきやすい点
    if cfg.common_pitfalls:
        heading(doc, "■ つまずきやすい点")
        for p in cfg.common_pitfalls:
            _bullet(doc, p)

    # 練習問題
    if cfg.problems:
        heading(doc, "■ 予習用練習問題")
        for i, q in enumerate(cfg.problems, 1):
            doc.add_paragraph(f"【問{i}】{q.question}")
            if q.math:
                add_math(doc, q.math)
            if q.answer:
                doc.add_paragraph(f"解答: {q.answer}")
            doc.add_paragraph()

    # まとめ
    if cfg.summary_paragraphs:
        heading(doc, "■ 授業前のまとめ")
        for p in cfg.summary_paragraphs:
            doc.add_paragraph(p)

    # 保存
    os.makedirs(output_dir, exist_ok=True)
    path = safe_save_path(os.path.join(output_dir, f"{filename_base}_v1.docx"))
    doc.save(path)
    return path


# === CLI として簡単テスト ===

if __name__ == "__main__":
    from omml_lib import m_det
    cfg = ReportConfig(
        subject="テスト科目",
        subtitle="サンプル：報告書の骨格テスト",
        videos=[{"id": "01_1", "title": "テスト動画", "duration": "10:00"}],
        goals=["骨格の確認"],
        sections=[
            Section(
                heading="§1 はじめに",
                paragraphs=["これはテスト段落。"],
                math=[MathBlock(caption="行列例:", inner_xml=m_det([[m_r('1'), m_r('2')], [m_r('3'), m_r('4')]]))],
            ),
        ],
        common_pitfalls=["テストの落とし穴"],
        # 標準構成は解答付き5〜6問（check_config が数える）
        problems=[Problem(question=f"{i}+1 = ?", answer=str(i + 1)) for i in range(1, 6)],
        summary_paragraphs=["以上。"],
    )
    BASE = os.environ.get("USERPROFILE") or str(Path.home())
    path = build_report(cfg, os.path.join(BASE, r"AppData\Local\Temp"), "report_gen_test")
    print("OK:", path)
